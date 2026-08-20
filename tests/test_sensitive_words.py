from __future__ import annotations

from app.services.sensitive_word_service import (
    apply_whitelist,
    keyword_hits_in_text,
    load_full_config,
    load_scan_profiles,
    normalize_text,
    save_full_config,
    scan_paths,
)
from app.json_io import write_json


def test_normalize_and_case():

    assert normalize_text("  Hello\n  World  ") == "Hello World"
    assert keyword_hits_in_text("Hello Battery Pack", "battery", [])
    assert not keyword_hits_in_text("Hello Pack", "battery", [])


def test_chinese_contains_and_multi_keyword():

    text = "儿童玩具电池套装"
    assert keyword_hits_in_text(text, "电池", [])
    assert keyword_hits_in_text(text, "儿童", [])
    assert keyword_hits_in_text("液 体硅胶", "液体", [])
    assert keyword_hits_in_text("液\u200b体硅胶", "液体", [])


def test_per_keyword_whitelist():

    text = "lithium battery for camera"
    assert keyword_hits_in_text(text, "battery", [])
    assert not keyword_hits_in_text(
        text,
        "battery",
        ["lithium battery"],
    )
    # whitelist of another keyword does not apply
    assert keyword_hits_in_text(
        text,
        "battery",
        ["camera"],
    )


def test_scan_profiles_default_and_start_row(tmp_path):

    from openpyxl import Workbook

    config = tmp_path / "sensitive_words.json"
    write_json(
        config,
        {
            "version": "1.1",
            "keywords": [
                {
                    "keyword": "电池",
                    "category": "合规",
                    "enabled": True,
                    "fields": ["ProductCN"],
                    "whitelist_phrases": [],
                }
            ],
        },
    )

    profiles = load_scan_profiles(config)
    assert profiles["MC"].data_start_row == 18
    assert profiles["KYD"].data_start_row == 30

    workbook = Workbook()
    ws = workbook.active
    ws.title = "迈创发票"
    ws["A2"] = "SKIP"
    ws["H2"] = "电池"
    ws["A17"] = "货箱编号*"
    ws["H17"] = "产品中文品名*"
    ws["A18"] = "FBA123U000001"
    ws["H18"] = "玩具电池"
    file_path = tmp_path / "迈创发票.xlsx"
    workbook.save(file_path)
    workbook.close()

    result = scan_paths([file_path], config_path=config)
    assert result.hit_count == 1
    assert result.hits[0].cell_address == "H18"


def test_load_old_config_merges_scan_profiles(tmp_path):

    config = tmp_path / "sensitive_words.json"
    write_json(
        config,
        {
            "version": "1.0",
            "keywords": [
                {
                    "keyword": "电池",
                    "category": "合规",
                    "enabled": True,
                    "fields": ["ProductCN"],
                    "whitelist_phrases": [],
                }
            ],
        },
    )

    data = load_full_config(config)
    codes = {item["carrier_code"] for item in data["scan_profiles"]}
    assert codes == {"MC", "KYD"}
    assert data["keywords"][0]["keyword"] == "电池"
    saved = load_full_config(config)
    assert saved["scan_profiles"]


def test_kyd_skips_rows_before_start(tmp_path):

    from openpyxl import Workbook

    config = tmp_path / "sensitive_words.json"
    write_json(
        config,
        {
            "version": "1.1",
            "keywords": [
                {
                    "keyword": "电池",
                    "category": "合规",
                    "enabled": True,
                    "fields": ["ProductCN"],
                    "whitelist_phrases": [],
                }
            ],
        },
    )

    workbook = Workbook()
    ws = workbook.active
    ws.title = "快越达发票"
    ws["A10"] = "FBA123U000001"
    ws["D10"] = "电池"
    ws["A29"] = "货箱编号*"
    ws["D29"] = "产品中文品名*"
    ws["A30"] = "FBA123U000002"
    ws["D30"] = "玩具电池"
    file_path = tmp_path / "快越达发票.xlsx"
    workbook.save(file_path)
    workbook.close()

    result = scan_paths([file_path], config_path=config)
    assert result.hit_count == 1
    assert result.hits[0].cell_address == "D30"


def test_custom_start_row_is_respected(tmp_path):

    from openpyxl import Workbook

    config = tmp_path / "sensitive_words.json"
    write_json(
        config,
        {
            "version": "1.1",
            "keywords": [
                {
                    "keyword": "电池",
                    "category": "合规",
                    "enabled": True,
                    "fields": ["ProductCN"],
                    "whitelist_phrases": [],
                }
            ],
            "scan_profiles": [
                {
                    "carrier_code": "MC",
                    "carrier_name": "迈创",
                    "data_start_row": 20,
                    "header_row": 19,
                    "sheet_name": "迈创发票",
                }
            ],
        },
    )

    profiles = load_scan_profiles(config)
    assert profiles["MC"].data_start_row == 20
    assert profiles["KYD"].data_start_row == 30

    workbook = Workbook()
    ws = workbook.active
    ws.title = "迈创发票"
    ws["A18"] = "FBA123U000001"
    ws["H18"] = "玩具电池"
    ws["A20"] = "FBA123U000002"
    ws["H20"] = "另一箱电池"
    file_path = tmp_path / "迈创发票.xlsx"
    workbook.save(file_path)
    workbook.close()

    result = scan_paths([file_path], config_path=config)
    assert result.hit_count == 1
    assert result.hits[0].cell_address == "H20"

    save_full_config(
        {
            "keywords": [],
            "scan_profiles": [
                {
                    "carrier_code": "MC",
                    "data_start_row": 22,
                }
            ],
        },
        config,
    )
    saved = load_scan_profiles(config)
    assert saved["MC"].data_start_row == 22


def test_field_scope_and_aggregation(tmp_path):

    from openpyxl import Workbook

    workbook = Workbook()
    ws = workbook.active
    ws.title = "Sheet1"
    ws.append(
        ["产品英文品名", "产品中文品名", "英文产品材质", "中文产品材质", "其他"]
    )
    ws.append(["Toy Battery", "玩具电池", "ABS", "塑料", "battery hidden"])
    file_path = tmp_path / "sample.xlsx"
    workbook.save(file_path)
    workbook.close()

    config = tmp_path / "sensitive_words.json"
    write_json(
        config,
        {
            "version": "1.0",
            "keywords": [
                {
                    "keyword": "battery",
                    "category": "合规",
                    "enabled": True,
                    "fields": ["ProductEN"],
                    "whitelist_phrases": [],
                },
                {
                    "keyword": "电池",
                    "category": "合规",
                    "enabled": True,
                    "fields": ["ProductCN"],
                    "whitelist_phrases": [],
                },
            ],
        },
    )

    result = scan_paths([file_path], config_path=config)

    assert result.hit_count == 2
    keywords = {hit.keyword for hit in result.hits}
    assert keywords == {"battery", "电池"}
    assert all(hit.field_code in {"ProductEN", "ProductCN"} for hit in result.hits)
    assert result.hit_file_count == 1
    assert len(result.aggregates) == 2
    assert all(item["cell_labels"] for item in result.aggregates)


def _liquid_config(tmp_path):

    config = tmp_path / "sensitive_words.json"
    write_json(
        config,
        {
            "version": "1.1",
            "keywords": [
                {
                    "keyword": "液体",
                    "category": "合规",
                    "enabled": True,
                    "fields": ["ProductCN"],
                    "whitelist_phrases": [],
                }
            ],
        },
    )
    return config


def test_product_cn_liquid_on_row_without_carton(tmp_path):

    from openpyxl import Workbook

    config = _liquid_config(tmp_path)
    workbook = Workbook()
    ws = workbook.active
    ws.title = "迈创发票"
    ws["E3"] = "液体*"
    ws["F3"] = "否"
    ws["A17"] = "货箱编号*"
    ws["H17"] = "产品中文品名*"
    ws["A18"] = "FBA123U000001"
    ws["H18"] = "雏菊项链"
    ws["H19"] = "液体硅胶玩具"
    file_path = tmp_path / "迈创发票.xlsx"
    workbook.save(file_path)
    workbook.close()

    result = scan_paths([file_path], config_path=config)
    assert result.hit_count == 1
    assert result.hits[0].cell_address == "H19"
    assert "液体硅胶玩具" in result.hits[0].original_text


def test_template_liquid_label_is_not_a_hit(tmp_path):

    from openpyxl import Workbook

    config = _liquid_config(tmp_path)
    workbook = Workbook()
    ws = workbook.active
    ws.title = "迈创发票"
    ws["E3"] = "液体*"
    ws["F3"] = "否"
    ws["A17"] = "货箱编号*"
    ws["H17"] = "产品中文品名*"
    ws["A18"] = "FBA123U000001"
    ws["H18"] = "雏菊项链"
    file_path = tmp_path / "迈创发票.xlsx"
    workbook.save(file_path)
    workbook.close()

    result = scan_paths([file_path], config_path=config)
    assert result.hit_count == 0


def test_product_cn_header_variant_and_invisible_chars(tmp_path):

    from openpyxl import Workbook

    config = _liquid_config(tmp_path)
    workbook = Workbook()
    ws = workbook.active
    ws.title = "迈创发票"
    ws["A17"] = "货箱编号*"
    ws["H17"] = "产品中文品名里"
    ws["A18"] = "FBA123U000001"
    ws["H18"] = "液\u200b体硅胶"
    file_path = tmp_path / "迈创发票.xlsx"
    workbook.save(file_path)
    workbook.close()

    result = scan_paths([file_path], config_path=config)
    assert result.hit_count == 1
    assert result.hits[0].cell_address == "H18"


def test_merged_invoice_sheet_hebing_hits_product_cn(tmp_path):

    from openpyxl import Workbook

    config = _liquid_config(tmp_path)
    workbook = Workbook()
    ws = workbook.active
    ws.title = "合并发票"
    ws["E3"] = "液体*"
    ws["F3"] = "否"
    ws["A17"] = "货箱编号*"
    ws["G17"] = "产品英文品名*"
    ws["H17"] = "产品中文品名*"
    ws["A18"] = "FBA19L0YW7PMU000001"
    ws["H18"] = "液体"
    file_path = tmp_path / "BJC1_迈创发票-BJC1-8箱-合并.xlsx"
    workbook.save(file_path)
    workbook.close()

    result = scan_paths([file_path], config_path=config)
    assert result.hit_count == 1
    assert result.hits[0].cell_address == "H18"
    assert result.hits[0].original_text == "液体"
    assert result.hits[0].sheet_name == "合并发票"


def test_format_cell_labels_and_suggested_value():

    from app.pages.sensitive_page import format_cell_labels, suggested_new_value

    assert format_cell_labels(["合并发票!H18"]) == "合并发票!H18"
    assert (
        format_cell_labels(
            ["合并发票!H18", "合并发票!H19", "合并发票!H20", "合并发票!H21"]
        )
        == "合并发票!H18、合并发票!H19、合并发票!H20 等4处"
    )
    assert suggested_new_value("液体", "液体") == ""
    assert suggested_new_value("液体硅胶玩具", "液体") == "硅胶玩具"
