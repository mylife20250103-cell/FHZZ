from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.services.address_library import (
    AddressRow,
    load_mc_address_rows,
    lookup_kyd_address,
)
from app.services.invoice_adapters.base import AdapterError


KYD_HEADERS = [
    "地址编码*",
    "FBA仓库代码",
    "联系人*",
    "公司名",
    "联系电话",
    "联系手机",
    "地址一*",
    "地址二",
    "地址三",
    "城市*",
    "省/洲",
    "国家*",
    "邮编*",
]

MC_HEADERS = [
    "地址简称*",
    "地址编码*",
    "FBA仓库代码",
    "联系人*",
    "公司名",
    "联系电话",
    "联系手机",
    "地址一*",
    "地址二",
    "地址三",
    "城市*",
    "省/洲",
    "国家*",
    "邮编*",
]


def _patch_central(monkeypatch, path: Path) -> None:
    import app.services.address_library as module

    monkeypatch.setattr(module, "CENTRAL_ADDRESS_PATH", path)


def test_lookup_kyd_matches_warehouse_not_first_row(
    tmp_path: Path,
    monkeypatch,
) -> None:
    path = tmp_path / "各物流地址库.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "快越达地址库"
    ws.append(KYD_HEADERS)
    ws.append(
        [
            "GGE1",
            "GGE1",
            "GGE1",
            "Amazon",
            111,
            None,
            "Old Street",
            None,
            None,
            "Foxborough",
            "MA",
            "US",
            "02035",
        ]
    )
    ws.append(
        [
            "GYR2",
            "GYR2",
            "GYR2",
            "Amazon",
            8553722005,
            None,
            "17341 W MINNEZONA AVE",
            None,
            None,
            "GOODYEAR",
            "AZ",
            "US",
            85395,
        ]
    )
    wb.save(path)

    _patch_central(monkeypatch, path)

    row = lookup_kyd_address("GYR2")
    assert isinstance(row, AddressRow)
    assert row.warehouse == "GYR2"
    assert row.city == "GOODYEAR"
    assert row.zip_code == 85395.0
    assert row.phone == 8553722005.0


def test_lookup_kyd_missing_warehouse(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "各物流地址库.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "快越达地址库"
    ws.append(KYD_HEADERS)
    ws.append(
        [
            "GGE1",
            "GGE1",
            "GGE1",
            "Amazon",
            111,
            None,
            "Old Street",
            None,
            None,
            "Foxborough",
            "MA",
            "US",
            "02035",
        ]
    )
    wb.save(path)
    _patch_central(monkeypatch, path)

    with pytest.raises(AdapterError, match="没有仓库"):
        lookup_kyd_address("GYR2")


def test_load_mc_rows_from_central_xlsx(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "各物流地址库.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "迈创地址库"
    ws.append(MC_HEADERS)
    ws.append(
        [
            "GYR2",
            "GYR2",
            "GYR2",
            "GYR2",
            "Amazon.com Services, Inc.",
            "1234567890",
            "1234567890",
            "17341 W MINNEZONA AVE",
            None,
            None,
            "GOODYEAR",
            "AZ",
            "US",
            "85395",
        ]
    )
    wb.save(path)
    _patch_central(monkeypatch, path)

    rows = load_mc_address_rows()
    assert rows[0][1] == "地址编码*"
    assert rows[1][1] == "GYR2"
    assert rows[1][10] == "GOODYEAR"


def test_mc_template_vlookup_uses_address_sheet() -> None:
    from app.invoice_config import MC_TEMPLATE_PATH

    if not MC_TEMPLATE_PATH.exists():
        return

    wb = load_workbook(MC_TEMPLATE_PATH)
    ws = wb["模板"]
    formula = str(ws["B6"].value)
    assert "VLOOKUP(B4,地址库!" in formula.replace(" ", "")
    assert "$N$990" in formula
    wb.close()
