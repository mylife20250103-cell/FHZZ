from __future__ import annotations

from pathlib import Path

from app.services.invoice_scan_service import (
    load_registered_sources,
    load_source_entries,
    resolve_invoice_scan_directories,
    scan_original_invoices,
)


def test_load_registered_sources_true(tmp_path, monkeypatch):

    ini_path = tmp_path / "invoice_sources.ini"

    ini_path.write_text(
        "\n".join(
            [
                "[Source.Mayn_KYD]",
                "Enabled=TRUE",
                r"Path=D:\invoices\kyd",
                "Owner=Mayn",
                "StoreCode=美10",
                "CarrierCode=KYD",
                "",
                "[Source.Mayn_MC]",
                "Enabled=TRUE",
                r"Path=D:\invoices\mc",
                "Owner=Mayn",
                "StoreCode=美10",
                "CarrierCode=MC",
                "",
                "[Source.Disabled]",
                "Enabled=FALSE",
                r"Path=D:\invoices\off",
            ]
        ),
        encoding="utf-8",
    )

    monkeypatch.setattr(
        "app.services.invoice_scan_service.INVOICE_SOURCES_INI",
        ini_path,
    )

    sources = load_registered_sources()

    assert sources == [
        Path(r"D:\invoices\kyd"),
        Path(r"D:\invoices\mc"),
    ]

    entries = load_source_entries()

    assert [entry.source_key for entry in entries] == [
        "Mayn_KYD",
        "Mayn_MC",
    ]
    assert entries[0].carrier_code == "KYD"
    assert entries[1].carrier_code == "MC"


def test_scan_does_not_depend_on_inquiry(tmp_path, monkeypatch):

    monkeypatch.setattr(
        "app.services.invoice_scan_service.load_known_carriers",
        lambda: {"KYD", "MC"},
    )

    result = scan_original_invoices(
        "20260819",
        [tmp_path],
    )

    joined = "\n".join(result.errors)

    assert "询价" not in joined
    assert "没有扫描到任何发票 .xlsx" in joined
    assert result.passed is False


def test_resolve_invoice_scan_directories_separates_source_and_extra():

    source = Path(r"D:\invoices\mc\未合并")
    extra = [
        Path(r"D:\invoices\mc\20260820"),
        Path(r"D:\invoices\kyd\20260820"),
    ]

    assert resolve_invoice_scan_directories(
        source_directory=source,
        extra_directories=extra,
        extra_only=False,
    ) == [source]

    assert resolve_invoice_scan_directories(
        source_directory=source,
        extra_directories=extra,
        extra_only=True,
    ) == extra

    assert resolve_invoice_scan_directories(
        source_directory=None,
        extra_directories=extra,
        extra_only=False,
    ) == []


def _write_meta_invoice(path: Path, **fields):

    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "快越达发票"
    ws = workbook.create_sheet("_SystemMeta")
    ws.sheet_state = "veryHidden"
    meta = {
        "MetaSchemaVersion": "1.0",
        "CarrierCode": "KYD",
        "CarrierName": "快越达",
        "TemplateVersion": "1.0",
        "InvoiceType": "KYD",
        "ChannelCell": "B4",
        "DateID": "20260819",
        "GeneratedAt": "2026-08-19 22:16:17",
        "StoreCode": "美10",
        "PlanID": "CD",
        "SourceID": "ECD6D9A4",
        "WarehouseCode": "GYR2",
        "FBABatch": "FBA3",
        "CartonNumber": "FBA3U000001",
    }
    meta.update(fields)
    for index, (key, value) in enumerate(meta.items(), start=1):
        ws.cell(index, 1, key)
        ws.cell(index, 2, value)
    workbook.save(path)
    workbook.close()


def test_scan_keeps_latest_generation_only(tmp_path, monkeypatch):

    monkeypatch.setattr(
        "app.services.invoice_scan_service.load_known_carriers",
        lambda: {"KYD", "MC"},
    )

    _write_meta_invoice(
        tmp_path / "old.xlsx",
        GeneratedAt="2026-08-19 22:16:17",
        CartonNumber="FBA3U000001",
        FBABatch="FBA3",
    )
    _write_meta_invoice(
        tmp_path / "new1.xlsx",
        GeneratedAt="2026-08-19 23:25:22",
        CartonNumber="FBA33U000001",
        FBABatch="FBA33",
        WarehouseCode="GYR2",
    )
    _write_meta_invoice(
        tmp_path / "new2.xlsx",
        GeneratedAt="2026-08-19 23:25:22",
        CartonNumber="FBA11U000001",
        FBABatch="FBA11",
        WarehouseCode="GYR3",
    )

    result = scan_original_invoices("20260819", [tmp_path])

    assert result.passed is True
    assert result.invoice_count == 2
    cartons = {record.carton_number for record in result.records}
    assert cartons == {"FBA33U000001", "FBA11U000001"}
    assert any("已跳过 1 个" in item for item in result.warnings)
    assert all("内部元数据不一致" not in item for item in result.errors)


def test_known_carriers_ignore_tracking_only_codes(tmp_path, monkeypatch):
    ini = tmp_path / "current_config.ini"
    ini.write_text(
        "\n".join(
            [
                "[Carrier.KYD]",
                "CarrierName=快越达",
                "[Carrier.MC]",
                "CarrierName=迈创",
                "[Carrier.HP]",
                "CarrierName=皓鹏",
                "[Carrier.LH]",
                "CarrierName=利合",
            ]
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(
        "app.services.invoice_scan_service.CURRENT_CONFIG_INI",
        ini,
    )
    from app.services.invoice_scan_service import load_known_carriers

    assert load_known_carriers() == {"KYD", "MC"}


def test_scan_includes_invoices_in_date_range(tmp_path, monkeypatch):

    monkeypatch.setattr(
        "app.services.invoice_scan_service.load_known_carriers",
        lambda: {"KYD", "MC"},
    )

    _write_meta_invoice(
        tmp_path / "day10.xlsx",
        DateID="20260910",
        SourceID="AAAAAAAA",
        CartonNumber="FBA10U000001",
        FBABatch="FBA10",
        GeneratedAt="2026-09-10 10:00:00",
    )
    _write_meta_invoice(
        tmp_path / "day11.xlsx",
        DateID="20260911",
        SourceID="BBBBBBBB",
        CartonNumber="FBA11U000001",
        FBABatch="FBA11",
        GeneratedAt="2026-09-11 10:00:00",
    )
    _write_meta_invoice(
        tmp_path / "day12.xlsx",
        DateID="20260912",
        SourceID="CCCCCCCC",
        CartonNumber="FBA12U000001",
        FBABatch="FBA12",
        GeneratedAt="2026-09-12 10:00:00",
    )

    single = scan_original_invoices("20260910", [tmp_path])
    assert {item.date_id for item in single.records} == {"20260910"}

    ranged = scan_original_invoices(
        "20260911",
        [tmp_path],
        allowed_date_ids=["20260910", "20260911"],
    )
    assert ranged.passed is True
    assert {item.date_id for item in ranged.records} == {
        "20260910",
        "20260911",
    }
    assert {item.source_id for item in ranged.records} == {
        "AAAAAAAA",
        "BBBBBBBB",
    }
