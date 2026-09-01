from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from app.services.batch_service import BatchRecord
from app.services.invoice_adapters.kyd_adapter import KydAdapter
from app.services.invoice_adapters.mc_adapter import McAdapter
from app.services.invoice_content_merge_service import _write_diagnose_report
from app.services.merge_diagnose_service import diagnose_merge_group


def _write_kyd_source(
    path: Path,
    *,
    carton="FBA19N000001",
    en="daisy necklace",
    cn="雏菊项链",
    unit=0.8,
    qty=70,
    hs=7113119090,
    url="",
    channel="K18美西稳速达-卡派-包税",
    warehouse="PSP3",
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "快越达发票"
    ws["B4"] = channel
    ws["B5"] = warehouse
    ws["A30"] = carton
    ws["C30"] = en
    ws["D30"] = cn
    ws["F30"] = unit
    ws["G30"] = qty
    ws["J30"] = hs
    ws["S30"] = url
    wb.save(path)
    wb.close()


def test_kyd_diagnose_source_reports_missing_hs_and_image(tmp_path):

    path = tmp_path / "PSP3_FBA19N000001_快越达发票.xlsx"
    _write_kyd_source(path, hs=None, url="")

    issues = KydAdapter().diagnose_source(path)

    assert any("海关编码为空" in item for item in issues)
    assert any("产品图片网址为空或不是网址" in item for item in issues)


def test_kyd_diagnose_source_reports_empty_names(tmp_path):

    path = tmp_path / "PSP3_FBA19N000002_快越达发票.xlsx"
    _write_kyd_source(path, en=None, cn="  ", url="")

    issues = KydAdapter().diagnose_source(path)

    assert any("产品英文品名为空" in item for item in issues)
    assert any("产品中文品名为空" in item for item in issues)


def test_mc_diagnose_source_reports_empty_product(tmp_path):

    path = tmp_path / "GYR2_FBA33U000001_迈创发票.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "迈创发票"
    ws["B3"] = "GYR2"
    ws["A18"] = "FBA33U000001"
    ws["G18"] = None
    ws["H18"] = "雏菊项链"
    wb.save(path)
    wb.close()

    issues = McAdapter().diagnose_source(path)

    assert any("产品英文品名为空" in item for item in issues)


def test_diagnose_merge_group_lists_each_file(tmp_path, monkeypatch):

    import app.services.merge_diagnose_service as module

    monkeypatch.setattr(module, "lookup_kyd_address", lambda warehouse: None)

    first = tmp_path / "PSP3_FBA19N000001_快越达发票.xlsx"
    second = tmp_path / "PSP3_FBA19N000002_快越达发票.xlsx"
    _write_kyd_source(first, hs=None, url="")
    _write_kyd_source(second, carton="FBA19N000002", en=None, url="")

    lines = diagnose_merge_group(
        KydAdapter(),
        {
            "warehouse_code": "PSP3",
            "input_count": 2,
            "input_files": [str(first), str(second)],
        },
    )

    text = "\n".join(lines)
    assert "【自检】KYD PSP3（2 箱）" in text
    assert "PSP3_FBA19N000001_快越达发票.xlsx" in text
    assert "海关编码为空" in text
    assert "产品英文品名为空" in text


def test_write_diagnose_report_saves_batch_file(tmp_path):

    batch = BatchRecord(
        batch_id="20260901-B0001",
        date_id="20260901",
        directory=tmp_path,
        status="merge_plan_ready",
        content_hash="",
        snapshot={},
        status_payload={},
    )
    path = Path(
        _write_diagnose_report(
            batch,
            ["【自检】KYD PSP3（1箱）", "- a.xlsx：海关编码为空"],
        )
    )

    assert path.name == "content_merge_diagnose.txt"
    text = path.read_text(encoding="utf-8")
    assert "20260901-B0001" in text
    assert "海关编码为空" in text
