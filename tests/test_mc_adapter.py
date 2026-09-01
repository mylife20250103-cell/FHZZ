from __future__ import annotations

from openpyxl import Workbook

from app.invoice_config import MC_TEMPLATE_PATH
from app.services.invoice_adapters.mc_adapter import McAdapter


def test_mc_template_anchors_confirmed():

    adapter = McAdapter()
    cfg = adapter.template_config()

    assert cfg.sheet_name == "迈创发票"
    assert adapter.OUTPUT_SHEET == "模板"
    assert adapter.output_extension() == ".xlsx"
    assert cfg.header_row == 17
    assert cfg.data_start_row == 18
    assert cfg.channel_cell == "B2"
    assert cfg.warehouse_cell == "B3"
    assert cfg.carton_count_cell == "B16"

    columns = {
        spec.field_code: spec.column
        for spec in cfg.product_fields
    }
    assert columns["ProductEN"] == 7
    assert columns["ProductCN"] == 8
    assert columns["MaterialEN"] == 11
    assert columns["MaterialCN"] == 11


def test_mc_validate_source_rejects_wrong_carrier(tmp_path):

    adapter = McAdapter()
    fake = tmp_path / "demo.xlsx"
    fake.write_bytes(b"not-excel")

    errors = adapter.validate_source(
        fake,
        {"CarrierCode": "KYD", "TemplateVersion": "1.0"},
    )

    assert any("CarrierCode 不是 MC" in item for item in errors)


def _write_source(path, warehouse: str, carton: str, url: str) -> None:

    wb = Workbook()
    ws = wb.active
    ws.title = "迈创发票"
    ws["B2"] = "墨西哥空运带电/敏感"
    ws["B3"] = warehouse
    ws["F1"] = "否"
    ws["A18"] = carton
    ws["B18"] = "C"
    ws["C18"] = 16
    ws["D18"] = 50
    ws["E18"] = 40
    ws["F18"] = 40
    ws["G18"] = "daisy necklace"
    ws["H18"] = "雏菊项链"
    ws["I18"] = 0.8
    ws["J18"] = 70
    ws["K18"] = "alloy"
    ws["L18"] = 7113119090
    ws["R18"] = url
    wb.save(path)
    wb.close()


def test_merge_group_copies_official_template(tmp_path):

    if not MC_TEMPLATE_PATH.exists():
        return

    source = tmp_path / "GYR2_FBA33U000001_迈创发票.xlsx"
    _write_source(
        source,
        "GYR2",
        "FBA33U000001",
        "https://example.com/demo.jpg",
    )
    template_mtime = MC_TEMPLATE_PATH.stat().st_mtime
    dest = tmp_path / "MC_GYR2_1箱.xlsx"
    adapter = McAdapter()
    group = {
        "batch_id": "TEST",
        "carrier_code": "MC",
        "template_version": "1.0",
        "warehouse_code": "GYR2",
        "input_count": 1,
        "input_files": [str(source)],
        "carton_numbers": ["FBA33U000001"],
        "source_ids": ["x"],
        "merge_plan_hash": "h",
        "date_id": "20260820",
    }
    adapter.merge_group(group, dest)

    assert dest.exists()
    assert MC_TEMPLATE_PATH.stat().st_mtime == template_mtime
    assert adapter.validate_output(dest, group) == []

    from openpyxl import load_workbook

    wb = load_workbook(dest, data_only=False)
    assert wb.sheetnames[:3] == ["模板", "渠道列表", "地址库"]
    ws = wb["模板"]
    assert ws["B2"].value == "美速达快提卡派"
    assert str(ws["B3"].value).strip() == "GYR2"
    assert isinstance(ws["B4"].value, str) and ws["B4"].value.startswith("=")
    assert int(float(ws["B16"].value)) == 1
    assert str(ws["A18"].value).strip() == "FBA33U000001"
    assert ws["R18"].value == "https://example.com/demo.jpg"
    assert "_MergeMeta" not in wb.sheetnames

    addr = wb["地址库"]
    codes = []
    for row in range(2, addr.max_row + 1):
        code = str(addr.cell(row, 2).value or "").strip().upper()
        if code:
            codes.append(code)
    assert "GYR2" in codes
    assert f"$N${len(codes) + 1}" in str(ws["B6"].value)
    wb.close()
