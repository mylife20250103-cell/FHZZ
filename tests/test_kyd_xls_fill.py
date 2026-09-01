from __future__ import annotations

from app.invoice_config import KYD_TEMPLATE_PATH
from app.services.invoice_adapters.kyd_adapter import KydAdapter
from app.services.kyd_xls_io import KydLine, copy_template, fill_kyd_template


# 最小合法 JPEG，仅用于结构测试。
MINI_JPEG = (
    b"\xff\xd8\xff\xe0\x00\x10JFIF\x00\x01\x01\x00\x00\x01\x00\x01\x00\x00"
    b"\xff\xdb\x00C\x00"
    + bytes([1] * 64)
    + b"\xff\xc0\x00\x0b\x08\x00\x01\x00\x01\x01\x01\x11\x00"
    b"\xff\xda\x00\x08\x01\x01\x00\x00?\x00\x7f\xff\xd9"
)


def test_kyd_adapter_keeps_source_sheet_and_xls_output():

    adapter = KydAdapter()
    cfg = adapter.template_config()
    assert cfg.sheet_name == "快越达发票"
    assert adapter.OUTPUT_SHEET == "模板"
    assert adapter.output_extension() == ".xls"


def test_fill_official_template_without_excel(tmp_path):

    if not KYD_TEMPLATE_PATH.exists():
        return

    dest = tmp_path / "KYD_GYR2_1箱.xls"
    copy_template(KYD_TEMPLATE_PATH, dest)

    fill_kyd_template(
        dest,
        channel="K18美西稳速达-卡派-包税",
        warehouse="GYR2",
        carton_count=1,
        lines=[
            KydLine(
                cells=[
                    "FBA33U000001",
                    "C",
                    "daisy necklace",
                    "雏菊项链",
                    1,
                    0.8,
                    70,
                    70,
                    56.0,
                    7113119090,
                    "无",
                    "无",
                    "alloy",
                    "日用",
                    16,
                    50,
                    40,
                    40,
                    None,
                    None,
                ],
                image_bytes=MINI_JPEG,
                image_kind="jpeg",
            )
        ],
    )

    import olefile
    import xlrd

    ole = olefile.OleFileIO(str(dest))
    streams = ["/".join(item) for item in ole.listdir()]
    assert not any("CompObj" in name for name in streams)
    book = ole.openstream(["Workbook"]).read()
    ole.close()

    wb = xlrd.open_workbook(str(dest))
    assert wb.sheet_names()[:3] == ["模板", "服务", "地址库"]
    ws = wb.sheet_by_name("模板")
    assert str(ws.cell_value(3, 1)).strip() == "K18美西稳速达-卡派-包税"
    assert str(ws.cell_value(4, 1)).strip() == "GYR2"
    assert str(ws.cell_value(5, 1)).strip() == "GYR2"
    assert int(float(ws.cell_value(11, 1))) == 85395
    assert int(float(ws.cell_value(25, 1))) == 1
    assert str(ws.cell_value(29, 0)).strip() == "FBA33U000001"
    assert ws.cell_value(29, 2) == "daisy necklace"

    addr = wb.sheet_by_name("地址库")
    codes = [str(addr.cell_value(row, 0)).strip().upper() for row in range(addr.nrows)]
    assert codes.count("BJC1") == 1
    assert codes.count("GYR2") == 1

    import struct
    from app.services.kyd_xls_io import iter_biff

    pics = 0
    indexes = 0
    dbcells = 0
    for _pos, rec, ln, payload in iter_biff(book):
        if rec == 0x020B:
            indexes += 1
        if rec == 0x00D7:
            dbcells += 1
        if rec == 0x005D and ln >= 8:
            _ft, _cb, ot, _oid = struct.unpack_from("<HHHH", payload, 0)
            if ot == 8:
                pics += 1
    assert pics >= 1
    assert indexes == 3
    assert dbcells >= 6
    recs = list(iter_biff(book))
    db_pos = [pos for pos, rec, _ln, _payload in recs if rec == 0x00D7]
    for _pos, rec, ln, payload in recs:
        if rec != 0x020B or ln < 20:
            continue
        count = (ln - 16) // 4
        pointers = struct.unpack_from("<" + "I" * count, payload, 16)
        for pointer in pointers:
            assert pointer in db_pos

    # 产品图必须算进工作表 DgContainer，否则快越达会直接报请求错误。
    bofs = [pos for pos, rec, ln, _payload in recs if rec == 0x0809]
    sheet_start, sheet_end = bofs[1], bofs[2]
    draws = []
    for pos, rec, ln, payload in recs:
        if pos < sheet_start or pos >= sheet_end:
            continue
        if rec == 0x00EC:
            draws.append(payload)
    concat = b"".join(draws)
    f002_ln = struct.unpack_from("<I", concat, 4)[0]
    assert f002_ln + 8 == len(concat)

    # 产品图必须是小尺寸渐进 JPEG，不能塞原图 800px。
    dg = bytearray()
    for i, (_pos, rec, ln, payload) in enumerate(recs):
        if rec == 0x00EB:
            dg.extend(payload)
            j = i + 1
            while j < len(recs) and recs[j][1] == 0x003C:
                dg.extend(recs[j][3])
                j += 1
            break
    soi = bytes(dg).find(b"\xff\xd8")
    eoi = bytes(dg).rfind(b"\xff\xd9")
    jpeg = bytes(dg)[soi : eoi + 2]
    assert 80 < len(jpeg) < 8000
    assert b"\xff\xc2" in jpeg

    # 快越达认的是 WPS 网址转图片那种 ID_ 浮动图（174 字节 F004），不是 HTCI_/单元格图包。
    id_pics = 0

    def count_id_pics(payload: bytes) -> None:
        nonlocal id_pics
        pos = 0
        while pos + 8 <= len(payload):
            verinst, typ, ln = struct.unpack_from("<HHI", payload, pos)
            ver = verinst & 0xF
            body = payload[pos + 8 : pos + 8 + min(ln, len(payload) - pos - 8)]
            if typ == 0xF004 and ver == 15 and ln == 174:
                if b"I\x00D\x00_" in body:
                    id_pics += 1
            if ver == 15:
                count_id_pics(body)
            pos += 8 + ln

    count_id_pics(concat)
    assert id_pics >= 1
    ole = olefile.OleFileIO(str(dest))
    streams = ["/".join(item) for item in ole.listdir()]
    ole.close()
    assert not any("ETCellImageData" in name for name in streams)


def _sample_photo_jpeg() -> bytes:
    from io import BytesIO

    from PIL import Image

    buf = BytesIO()
    Image.new("RGB", (220, 220), (40, 80, 160)).save(buf, "JPEG", quality=92)
    return buf.getvalue()


def _one_line(index: int, image: bytes) -> KydLine:
    return KydLine(
        cells=[
            f"FBA33U{index:06d}",
            "C",
            "daisy necklace",
            "雏菊项链",
            1,
            0.8,
            70,
            70,
            56.0,
            7113119090,
            "无",
            "无",
            "alloy",
            "日用",
            16,
            50,
            40,
            40,
            None,
            None,
        ],
        image_bytes=image,
        image_kind="jpeg",
    )


def test_ten_box_blips_stay_in_first_drawing_record(tmp_path):

    if not KYD_TEMPLATE_PATH.exists():
        return

    dest = tmp_path / "KYD_PSP3_10箱.xls"
    copy_template(KYD_TEMPLATE_PATH, dest)
    photo = _sample_photo_jpeg()
    fill_kyd_template(
        dest,
        channel="K18美西稳速达-卡派-包税",
        warehouse="GYR2",
        carton_count=10,
        lines=[_one_line(i, photo) for i in range(10)],
    )

    import olefile
    import struct
    from app.services.kyd_xls_io import iter_biff

    ole = olefile.OleFileIO(str(dest))
    book = ole.openstream(["Workbook"]).read()
    ole.close()
    recs = list(iter_biff(book))
    chunks = []
    for i, (_pos, rec, ln, payload) in enumerate(recs):
        if rec != 0x00EB:
            continue
        chunks.append(payload)
        j = i + 1
        while j < len(recs) and recs[j][1] == 0x003C:
            chunks.append(recs[j][3])
            j += 1
        break
    assert chunks
    first_len = len(chunks[0])
    whole = b"".join(chunks)
    _vi, _typ, ln = struct.unpack_from("<HHI", whole, 0)
    inner = whole[8 : 8 + ln]
    pos = 0
    complete_in_continue = 0
    while pos + 8 <= len(inner):
        cvi, ctyp, cln = struct.unpack_from("<HHI", inner, pos)
        if ctyp == 0xF001:
            body_off = 8 + pos + 8
            p = 0
            while p + 8 <= cln:
                _bvi, btyp, bln = struct.unpack_from("<HHI", inner, pos + 8 + p)
                abs_off = body_off + p
                end = abs_off + 8 + bln
                if btyp == 0xF007 and abs_off >= first_len:
                    complete_in_continue += 1
                p += 8 + bln
        pos += 8 + cln
    assert complete_in_continue == 0


def test_ten_box_writes_total_value_as_number_not_shared_formula(tmp_path):

    if not KYD_TEMPLATE_PATH.exists():
        return

    dest = tmp_path / "KYD_PSP3_10箱.xls"
    copy_template(KYD_TEMPLATE_PATH, dest)
    photo = _sample_photo_jpeg()
    lines = []
    for index in range(10):
        line = _one_line(index, photo)
        line.cells[3] = "钥匙扣磁铁" if index == 8 else "玻璃罐密封圈白色"
        line.cells[8] = 0.8 * (100 + index)
        line.cells[9] = 3926400000 + index
        lines.append(line)
    fill_kyd_template(
        dest,
        channel="K18美西稳速达-卡派-包税",
        warehouse="GYR2",
        carton_count=10,
        lines=lines,
    )

    import olefile
    import struct
    from app.services.kyd_xls_io import iter_biff, template_sheet_span

    ole = olefile.OleFileIO(str(dest))
    book = ole.openstream(["Workbook"]).read()
    ole.close()
    start, end = template_sheet_span(book)
    formulas = 0
    shared = 0
    totals = {}
    hs_codes = {}
    for pos, rec, ln, payload in iter_biff(book):
        if pos < start or pos >= end:
            continue
        if rec == 0x04BC:
            shared += 1
        if rec == 0x0006 and len(payload) >= 6:
            row, col, _xf = struct.unpack_from("<HHH", payload, 0)
            if col == 8 and row >= 29:
                formulas += 1
        if rec == 0x0203 and len(payload) >= 14:
            row, col, _xf = struct.unpack_from("<HHH", payload, 0)
            value = struct.unpack_from("<d", payload, 6)[0]
            if col == 8 and 29 <= row <= 38:
                totals[row] = value
            if col == 9 and 29 <= row <= 38:
                hs_codes[row] = value
    assert formulas == 0
    assert shared == 0
    assert len(totals) == 10
    assert hs_codes[37] == 3926400008
    assert abs(totals[38] - 0.8 * 109) < 1e-6


def test_merge_group_copies_central_template(tmp_path, monkeypatch):

    if not KYD_TEMPLATE_PATH.exists():
        return

    from openpyxl import Workbook
    import app.services.invoice_adapters.kyd_adapter as ka

    source = tmp_path / "GYR2_FBA33U000001.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "快越达发票"
    ws["B4"] = "K18美西稳速达-卡派-包税"
    ws["B5"] = "GYR2"
    ws["A30"] = "FBA33U000001"
    ws["C30"] = "daisy necklace"
    ws["D30"] = "雏菊项链"
    ws["F30"] = 0.8
    ws["G30"] = 70
    ws["S30"] = "https://example.com/demo.jpg"
    wb.save(source)
    wb.close()

    monkeypatch.setattr(
        ka,
        "download_product_image",
        lambda url, name: (MINI_JPEG, "jpeg"),
    )

    dest = tmp_path / "KYD_GYR2_1箱.xls"
    adapter = KydAdapter()
    group = {
        "batch_id": "TEST",
        "carrier_code": "KYD",
        "template_version": "1.0",
        "warehouse_code": "GYR2",
        "input_count": 1,
        "input_files": [str(source)],
        "carton_numbers": ["FBA33U000001"],
        "source_ids": ["x"],
        "merge_plan_hash": "h",
        "date_id": "20260820",
    }
    adapter.merge_group(group, dest)
    assert dest.exists()
    assert adapter.validate_output(dest, group) == []


def _filled_kyd_xls(dest):
    copy_template(KYD_TEMPLATE_PATH, dest)
    fill_kyd_template(
        dest,
        channel="K18美西稳速达-卡派-包税",
        warehouse="GYR2",
        carton_count=1,
        lines=[
            KydLine(
                cells=[
                    "FBA33U000001",
                    "C",
                    "daisy necklace",
                    "雏菊项链",
                    1,
                    0.8,
                    70,
                    70,
                    56.0,
                    7113119090,
                    "无",
                    "无",
                    "alloy",
                    "日用",
                    16,
                    50,
                    40,
                    40,
                    None,
                    None,
                ],
                image_bytes=MINI_JPEG,
                image_kind="jpeg",
            )
        ],
    )


def test_set_kyd_text_cell_does_not_add_compobj(tmp_path):

    if not KYD_TEMPLATE_PATH.exists():
        return

    import olefile
    import xlrd

    from app.services.kyd_xls_io import set_kyd_text_cell

    dest = tmp_path / "KYD_GYR2_1箱.xls"
    _filled_kyd_xls(dest)
    set_kyd_text_cell(dest, "模板", 3, 1, "K18美西顺丰速运-卡派")

    ole = olefile.OleFileIO(str(dest))
    streams = ["/".join(item) for item in ole.listdir()]
    book = ole.openstream(["Workbook"]).read()
    ole.close()
    assert not any("CompObj" in name for name in streams)

    pics = 0
    import struct
    from app.services.kyd_xls_io import iter_biff

    for _pos, rec, ln, payload in iter_biff(book):
        if rec == 0x005D and ln >= 8:
            _ft, _cb, ot, _oid = struct.unpack_from("<HHHH", payload, 0)
            if ot == 8:
                pics += 1
    assert pics >= 1

    wb = xlrd.open_workbook(str(dest))
    ws = wb.sheet_by_name("模板")
    assert str(ws.cell_value(3, 1)).strip() == "K18美西顺丰速运-卡派"
    assert str(ws.cell_value(4, 1)).strip() == "GYR2"
    assert str(ws.cell_value(29, 0)).strip() == "FBA33U000001"


def test_prefer_jpeg_url_strips_amazon_fmwebp():

    from app.services.invoice_adapters.kyd_adapter import prefer_jpeg_url

    assert prefer_jpeg_url(
        "https://m.media-amazon.com/images/I/6120IUo3pxL._AC_UL640_FMwebp_QL65_.jpg"
    ) == "https://m.media-amazon.com/images/I/6120IUo3pxL._AC_UL640_QL65_.jpg"
    assert prefer_jpeg_url(
        "https://m.media-amazon.com/images/I/51IDZfy2nqL.__AC_SX300_SY300_QL70_FMwebp_.jpg"
    ) == "https://m.media-amazon.com/images/I/51IDZfy2nqL.__AC_SX300_SY300_QL70_.jpg"


def test_detect_image_kind_accepts_webp_and_prepare_converts():

    from io import BytesIO

    from PIL import Image

    from app.services.invoice_adapters.kyd_adapter import detect_image_kind
    from app.services.kyd_xls_io import prepare_kyd_picture

    buf = BytesIO()
    Image.new("RGB", (40, 40), "red").save(buf, "WEBP")
    webp = buf.getvalue()
    assert detect_image_kind(webp) == "webp"
    jpeg = prepare_kyd_picture(webp)
    assert jpeg.startswith(b"\xff\xd8")
    small = prepare_kyd_picture(_sample_photo_jpeg(), max_bytes=600)
    assert 80 <= len(small) <= 600


def test_download_product_image_tries_jpeg_url_first(monkeypatch):

    from app.services.invoice_adapters import kyd_adapter as ka

    calls = []

    def fake_get(url: str) -> bytes:
        calls.append(url)
        if "FMwebp" in url:
            return b"RIFF\x00\x00\x00\x00WEBPVP8 "
        return MINI_JPEG

    monkeypatch.setattr(ka, "_http_get_image", fake_get)
    data, kind = ka.download_product_image(
        "https://m.media-amazon.com/images/I/abc._AC_UL640_FMwebp_QL65_.jpg",
        "demo.xlsx",
    )
    assert kind == "jpeg"
    assert data.startswith(b"\xff\xd8")
    assert "FMwebp" not in calls[0]


