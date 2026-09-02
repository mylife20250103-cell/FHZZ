from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook

from app.shipment_tracking.indexing.shipping_file_indexer import IndexedShippingFile
from app.shipment_tracking.services.shipping_load_service import (
    load_shipping_snapshot,
    load_store_snapshots,
)


def _write_plan(path: Path, fba_id: str, sku: str, qty: int) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "美国3号-发货规划"
    ws["M1"] = "仓库"
    ws["M2"] = "FBA批号"
    ws["N1"] = "IND9"
    ws["N2"] = fba_id
    ws["N13"] = 1
    ws["A13"] = "*产品英文品名"
    ws["F13"] = "SKU"
    ws["F14"] = sku
    ws["N14"] = qty
    wb.save(path)
    wb.close()


def _indexed(path: Path) -> IndexedShippingFile:
    return IndexedShippingFile(
        path=path,
        store_code="美3",
        year=2026,
        month=6,
        folder_day=11,
        ship_date=date(2026, 6, 11),
        filename_month=6,
        filename_day=11,
        batch_no=1,
        candidate_forwarder="迈创合德",
    )


def test_load_snapshot_is_batch_fba_sku_qty(tmp_path):
    path = tmp_path / "0611-美3-迈创合德.xlsx"
    _write_plan(path, "FBA-A", "SKU-A", 40)
    snap = load_shipping_snapshot(_indexed(path))

    assert snap.source.store_code == "美3"
    assert snap.source.batch_no == 1
    assert snap.source.ship_date == date(2026, 6, 11)
    assert snap.plan.fbas[0].fba_id == "FBA-A"
    assert snap.plan.fbas[0].cartons[0].items[0].sku == "SKU-A"
    assert snap.aggregated[0].items[0].quantity == 40


def test_load_store_skips_unreadable_file(tmp_path):
    root = tmp_path / "装箱明细"
    good = root / "美3" / "2026.6" / "6.11美3" / "0611-美3-迈创合德.xlsx"
    bad = root / "美3" / "2026.6" / "6.11美3" / "0611-美3-坏文件.xlsx"
    _write_plan(good, "FBA-A", "SKU-A", 8)
    bad.parent.mkdir(parents=True, exist_ok=True)
    bad.write_bytes(b"not-excel")

    result = load_store_snapshots(
        root,
        "美3",
        batch_root=tmp_path / "batches",
        merge_root=tmp_path / "merges",
    )
    assert len(result.snapshots) == 1
    assert result.snapshots[0].aggregated[0].items[0].quantity == 8
    assert any("坏文件" in item for item in result.errors)
    assert result.forwarders == {}
    assert result.channels == {}


def test_load_store_attaches_confirmed_forwarder(tmp_path):
    from app.json_io import write_json

    root = tmp_path / "装箱明细"
    plan = root / "美3" / "2026.6" / "6.11美3" / "0611-美3-迈创合德.xlsx"
    _write_plan(plan, "FBA-A", "SKU-A", 8)
    batch_dir = tmp_path / "batches" / "20260611" / "20260611-B0001"
    batch_dir.mkdir(parents=True)
    write_json(
        batch_dir / "snapshot.json",
        {
            "Files": [
                {
                    "store_code": "美3",
                    "fba_batch": "FBA-A",
                    "carrier_code": "KYD",
                }
            ]
        },
    )
    write_json(
        batch_dir / "status.json",
        {"Status": "FIRST_SCAN_PASSED"},
    )

    result = load_store_snapshots(
        root,
        "美3",
        batch_root=tmp_path / "batches",
        merge_root=tmp_path / "merges",
    )
    bound = result.forwarders["FBA-A"]
    assert bound.state == "confirmed"
    assert bound.carrier_code == "KYD"


def test_load_store_can_limit_to_one_month(tmp_path):
    root = tmp_path / "装箱明细"
    june = root / "美3" / "2026.6" / "6.11美3" / "0611-美3-迈创合德.xlsx"
    august = root / "美3" / "2026.8" / "8.21美3" / "0821-美3-快越达.xlsx"
    _write_plan(june, "FBA-A", "SKU-A", 8)
    _write_plan(august, "FBA-B", "SKU-B", 3)

    june_only = load_store_snapshots(
        root,
        "美3",
        batch_root=tmp_path / "batches",
        merge_root=tmp_path / "merges",
        year=2026,
        month=6,
        all_periods=False,
    )
    assert len(june_only.snapshots) == 1
    assert june_only.snapshots[0].aggregated[0].fba_id == "FBA-A"

    all_months = load_store_snapshots(
        root,
        "美3",
        batch_root=tmp_path / "batches",
        merge_root=tmp_path / "merges",
        all_periods=True,
    )
    assert len(all_months.snapshots) == 2


def test_load_store_can_limit_to_date_range(tmp_path):
    root = tmp_path / "装箱明细"
    june = root / "美3" / "2026.6" / "6.11美3" / "0611-美3-迈创合德.xlsx"
    august = root / "美3" / "2026.8" / "8.21美3" / "0821-美3-快越达.xlsx"
    _write_plan(june, "FBA-A", "SKU-A", 8)
    _write_plan(august, "FBA-B", "SKU-B", 3)

    ranged = load_store_snapshots(
        root,
        "美3",
        batch_root=tmp_path / "batches",
        merge_root=tmp_path / "merges",
        start_date=date(2026, 8, 1),
        end_date=date(2026, 8, 31),
    )
    assert len(ranged.snapshots) == 1
    assert ranged.snapshots[0].aggregated[0].fba_id == "FBA-B"


def test_load_store_can_scan_all_stores_in_date_range(tmp_path):
    root = tmp_path / "装箱明细"
    store_three = root / "美3" / "2026.8" / "8.21美3" / "0821-美3-快越达.xlsx"
    store_ten = root / "美10" / "2026.8" / "8.19美10" / "0819-美10-迈创合德.xlsx"
    other_month = root / "美3" / "2026.6" / "6.11美3" / "0611-美3-迈创合德.xlsx"
    _write_plan(store_three, "FBA-A", "SKU-A", 8)
    _write_plan(store_ten, "FBA-B", "SKU-B", 3)
    _write_plan(other_month, "FBA-C", "SKU-C", 5)

    ranged = load_store_snapshots(
        root,
        None,
        batch_root=tmp_path / "batches",
        merge_root=tmp_path / "merges",
        start_date=date(2026, 8, 1),
        end_date=date(2026, 8, 31),
    )
    fbas = {snap.aggregated[0].fba_id for snap in ranged.snapshots}
    stores = {snap.source.store_code for snap in ranged.snapshots}
    assert fbas == {"FBA-A", "FBA-B"}
    assert stores == {"美3", "美10"}
