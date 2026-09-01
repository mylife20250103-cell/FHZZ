from __future__ import annotations

from app.shipment_tracking.aggregation.fba_aggregator import aggregate_fba_sku_qty
from app.shipment_tracking.domain.parsed_plan import (
    ParsedCarton,
    ParsedCartonItem,
    ParsedFBA,
    ParsedTrackingPlan,
)
from app.shipment_tracking.parsing.shipping_plan_reader import parse_shipping_plan


def _carton(
    fba_id: str,
    carton_no: str,
    items: list[ParsedCartonItem],
    warehouse: str | None = "IND9",
) -> ParsedCarton:
    return ParsedCarton(
        fba_id=fba_id,
        destination_fc=warehouse,
        carton_no=carton_no,
        weight_kg=None,
        length_cm=None,
        width_cm=None,
        height_cm=None,
        items=tuple(items),
    )


def test_aggregate_one_fba_multi_sku():
    plan = ParsedTrackingPlan(
        fbas=(
            ParsedFBA(
                fba_id="FBA-A",
                destination_fc="IND9",
                cartons=(
                    _carton(
                        "FBA-A",
                        "1",
                        [
                            ParsedCartonItem(sku="SKU-1", quantity=8),
                            ParsedCartonItem(sku="SKU-2", quantity=50),
                            ParsedCartonItem(sku="SKU-3", quantity=20),
                        ],
                    ),
                ),
            ),
        )
    )
    fbas = aggregate_fba_sku_qty(plan)
    assert len(fbas) == 1
    qty = {item.sku: item.quantity for item in fbas[0].items}
    assert qty == {"SKU-1": 8, "SKU-2": 50, "SKU-3": 20}


def test_aggregate_sku_across_cartons_same_fba():
    plan = ParsedTrackingPlan(
        fbas=(
            ParsedFBA(
                fba_id="FBA-A",
                destination_fc="IND9",
                cartons=(
                    _carton(
                        "FBA-A",
                        "1",
                        [ParsedCartonItem(sku="SKU-A", quantity=20)],
                    ),
                    _carton(
                        "FBA-A",
                        "2",
                        [ParsedCartonItem(sku="SKU-A", quantity=30)],
                    ),
                ),
            ),
        )
    )
    fbas = aggregate_fba_sku_qty(plan)
    assert len(fbas) == 1
    assert fbas[0].items[0].sku == "SKU-A"
    assert fbas[0].items[0].quantity == 50


def test_aggregate_same_sku_two_fbas_not_merged():
    plan = ParsedTrackingPlan(
        fbas=(
            ParsedFBA(
                fba_id="FBA-A",
                destination_fc="IND9",
                cartons=(
                    _carton(
                        "FBA-A",
                        "1",
                        [ParsedCartonItem(sku="SKU-A", quantity=100)],
                    ),
                ),
            ),
            ParsedFBA(
                fba_id="FBA-B",
                destination_fc="ABE8",
                cartons=(
                    _carton(
                        "FBA-B",
                        "2",
                        [ParsedCartonItem(sku="SKU-A", quantity=200)],
                        warehouse="ABE8",
                    ),
                ),
            ),
        )
    )
    fbas = aggregate_fba_sku_qty(plan)
    assert [item.fba_id for item in fbas] == ["FBA-A", "FBA-B"]
    assert fbas[0].items[0].quantity == 100
    assert fbas[1].items[0].quantity == 200
    assert fbas[0].items[0].quantity + fbas[1].items[0].quantity == 300


def test_aggregate_after_parse(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "plan.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "美国3号-发货规划"
    ws["M1"] = "仓库"
    ws["M2"] = "FBA批号"
    ws["N1"] = "IND9"
    ws["N2"] = "FBA19AAA"
    ws["N13"] = 1
    ws["A13"] = "*产品英文品名"
    ws["F13"] = "SKU"
    ws["F14"] = "SKU-A"
    ws["N14"] = 20
    ws["F15"] = "SKU-B"
    ws["N15"] = 10
    wb.save(path)
    wb.close()

    plan = parse_shipping_plan(path, store_code="美3")
    fbas = aggregate_fba_sku_qty(plan)
    qty = {item.sku: item.quantity for item in fbas[0].items}
    assert qty == {"SKU-A": 20, "SKU-B": 10}
