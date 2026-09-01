from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from app.shipment_tracking.parsing.shipping_plan_reader import parse_shipping_plan


def _write_plan(path: Path, columns: list[dict], sku_rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "美国3号-发货规划"
    ws["A1"] = "美3"
    ws["M1"] = "仓库"
    ws["M2"] = "FBA批号"
    ws["M3"] = "内部编码ID"
    ws["M4"] = "长(cm)"
    ws["M5"] = "宽(cm)"
    ws["M6"] = "高(cm)"
    ws["M7"] = "重量(KG)"
    ws["M12"] = "总数量(箱)"
    ws["A13"] = "*产品英文品名"
    ws["E13"] = "子ASIN"
    ws["F13"] = "SKU"
    ws["H13"] = "产品中文品名"
    ws["I13"] = "是否发货"
    ws["M13"] = "发货数量"

    for index, column in enumerate(columns):
        col = 14 + index
        ws.cell(1, col, column["warehouse"])
        ws.cell(2, col, column["fba_id"])
        ws.cell(4, col, column.get("length_cm", 40))
        ws.cell(5, col, column.get("width_cm", 40))
        ws.cell(6, col, column.get("height_cm", 40))
        ws.cell(7, col, column.get("weight_kg", 12.5))
        ws.cell(12, col, column.get("carton_count", 1))
        ws.cell(13, col, column.get("carton_no", index + 1))

    for offset, row in enumerate(sku_rows):
        excel_row = 14 + offset
        ws.cell(excel_row, 1, row.get("en"))
        ws.cell(excel_row, 5, row.get("asin"))
        ws.cell(excel_row, 6, row["sku"])
        ws.cell(excel_row, 8, row.get("cn"))
        ws.cell(excel_row, 9, row.get("ship", 1))
        qtys = row["qtys"]
        ws.cell(excel_row, 13, sum(qtys))
        for index, qty in enumerate(qtys):
            if qty:
                ws.cell(excel_row, 14 + index, qty)

    wb.create_sheet("美国8号-发货规划 1")
    wb.save(path)
    wb.close()


def _fba_sku_qty(plan, fba_id: str, sku: str) -> int:
    total = 0
    for fba in plan.fbas:
        if fba.fba_id != fba_id:
            continue
        for carton in fba.cartons:
            for item in carton.items:
                if item.sku == sku:
                    total += item.quantity
    return total


def test_parse_keeps_fba_carton_sku_qty(tmp_path):
    path = tmp_path / "plan.xlsx"
    _write_plan(
        path,
        columns=[
            {"warehouse": "IND9", "fba_id": "FBA19AAA", "carton_no": 1},
        ],
        sku_rows=[
            {
                "sku": "SKU-A",
                "asin": "B0A",
                "en": "alpha",
                "cn": "甲",
                "qtys": [20],
            },
            {
                "sku": "SKU-B",
                "asin": "B0B",
                "en": "beta",
                "cn": "乙",
                "qtys": [10],
            },
        ],
    )

    plan = parse_shipping_plan(path, store_code="美3")
    assert len(plan.fbas) == 1
    fba = plan.fbas[0]
    assert fba.fba_id == "FBA19AAA"
    assert fba.destination_fc == "IND9"
    assert len(fba.cartons) == 1
    carton = fba.cartons[0]
    assert carton.carton_no == "1"
    assert carton.fba_id == "FBA19AAA"
    skus = {item.sku: item.quantity for item in carton.items}
    assert skus == {"SKU-A": 20, "SKU-B": 10}


def test_one_fba_one_sku(tmp_path):
    path = tmp_path / "one.xlsx"
    _write_plan(
        path,
        columns=[{"warehouse": "ABE8", "fba_id": "FBA-A", "carton_no": 1}],
        sku_rows=[{"sku": "SKU-1", "qtys": [18]}],
    )
    plan = parse_shipping_plan(path, store_code="美3")
    assert [fba.fba_id for fba in plan.fbas] == ["FBA-A"]
    assert _fba_sku_qty(plan, "FBA-A", "SKU-1") == 18


def test_one_fba_multi_sku(tmp_path):
    path = tmp_path / "multi-sku.xlsx"
    _write_plan(
        path,
        columns=[{"warehouse": "PSP3", "fba_id": "FBA-A", "carton_no": 1}],
        sku_rows=[
            {"sku": "SKU-1", "qtys": [8]},
            {"sku": "SKU-2", "qtys": [50]},
            {"sku": "SKU-3", "qtys": [20]},
        ],
    )
    plan = parse_shipping_plan(path, store_code="美3")
    assert _fba_sku_qty(plan, "FBA-A", "SKU-1") == 8
    assert _fba_sku_qty(plan, "FBA-A", "SKU-2") == 50
    assert _fba_sku_qty(plan, "FBA-A", "SKU-3") == 20
    assert len(plan.fbas[0].cartons[0].items) == 3


def test_sku_across_cartons_same_fba(tmp_path):
    path = tmp_path / "split.xlsx"
    _write_plan(
        path,
        columns=[
            {"warehouse": "IND9", "fba_id": "FBA-A", "carton_no": 1},
            {"warehouse": "IND9", "fba_id": "FBA-A", "carton_no": 2},
        ],
        sku_rows=[{"sku": "SKU-A", "qtys": [20, 30]}],
    )
    plan = parse_shipping_plan(path, store_code="美3")
    assert len(plan.fbas) == 1
    cartons = plan.fbas[0].cartons
    assert len(cartons) == 2
    assert [c.carton_no for c in cartons] == ["1", "2"]
    assert [item.quantity for item in cartons[0].items] == [20]
    assert [item.quantity for item in cartons[1].items] == [30]
    assert _fba_sku_qty(plan, "FBA-A", "SKU-A") == 50


def test_same_sku_two_fbas_not_merged(tmp_path):
    path = tmp_path / "two-fba.xlsx"
    _write_plan(
        path,
        columns=[
            {"warehouse": "IND9", "fba_id": "FBA-A", "carton_no": 1},
            {"warehouse": "ABE8", "fba_id": "FBA-B", "carton_no": 2},
        ],
        sku_rows=[{"sku": "SKU-A", "qtys": [100, 200]}],
    )
    plan = parse_shipping_plan(path, store_code="美3")
    assert [fba.fba_id for fba in plan.fbas] == ["FBA-A", "FBA-B"]
    assert _fba_sku_qty(plan, "FBA-A", "SKU-A") == 100
    assert _fba_sku_qty(plan, "FBA-B", "SKU-A") == 200
    assert _fba_sku_qty(plan, "FBA-A", "SKU-A") + _fba_sku_qty(
        plan, "FBA-B", "SKU-A"
    ) == 300


def test_one_carton_multi_sku(tmp_path):
    path = tmp_path / "carton-multi.xlsx"
    _write_plan(
        path,
        columns=[{"warehouse": "SBD1", "fba_id": "FBA-A", "carton_no": 1}],
        sku_rows=[
            {"sku": "SKU-A", "qtys": [20]},
            {"sku": "SKU-B", "qtys": [10]},
        ],
    )
    carton = parse_shipping_plan(path, store_code="美3").fbas[0].cartons[0]
    assert [item.sku for item in carton.items] == ["SKU-A", "SKU-B"]
    assert carton.items[0].quantity == 20
    assert carton.items[1].quantity == 10


def test_skips_zero_qty_and_residual_sheet(tmp_path):
    path = tmp_path / "zeros.xlsx"
    _write_plan(
        path,
        columns=[{"warehouse": "SWF2", "fba_id": "FBA-A", "carton_no": 1}],
        sku_rows=[
            {"sku": "SKIP", "qtys": [0]},
            {"sku": "KEEP", "qtys": [8]},
        ],
    )
    plan = parse_shipping_plan(path, store_code="美3")
    skus = [item.sku for item in plan.fbas[0].cartons[0].items]
    assert skus == ["KEEP"]


def test_real_xlsm_multi_sku_and_same_sku_across_fbas():
    path = (
        Path(r"I:\OneDrive\OneDrive - Lion")
        / "【采购仓储】"
        / "【装箱明细】"
        / "美3"
        / "2026.6"
        / "6.11美3"
        / "1"
        / "0611-美3-迈创合德.xlsm"
    )
    if not path.exists():
        return

    plan = parse_shipping_plan(path, store_code="美3")
    ids = [fba.fba_id for fba in plan.fbas]
    assert ids == [
        "FBA19FY7J9B7",
        "FBA19FY62KTY",
        "FBA19FY55P7C",
        "FBA19FY7GY5S",
        "FBA19FY6DPXQ",
    ]
    first = plan.fbas[0]
    assert first.destination_fc == "IND9"
    skus = {item.sku: item.quantity for item in first.cartons[0].items}
    assert skus["US3A16-Silver"] == 8
    assert skus["US3A26-Orange"] == 50
    assert len(skus) == 5
    assert _fba_sku_qty(plan, "FBA19FY7J9B7", "US3A16-Silver") == 8
    assert _fba_sku_qty(plan, "FBA19FY62KTY", "US3A16-Silver") == 8
    assert _fba_sku_qty(plan, "FBA19FY7J9B7", "US3A16-Silver") != 40

