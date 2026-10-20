from pathlib import Path

import pytest
from openpyxl import Workbook

from app.json_io import read_json, write_json
from app.services import address_library, batch_service, invoice_scan_service
from app.services.invoice_scan_service import scan_original_invoices
from tests.test_address_library import KYD_HEADERS, MC_HEADERS
from tests.test_batch_service import make_record, make_result
from tests.test_invoice_sources import _write_meta_invoice


@pytest.fixture
def country_address_book(tmp_path_factory, monkeypatch):
    path = tmp_path_factory.mktemp("country-addresses") / "addresses.xlsx"
    wb = Workbook()
    for name, headers, code_col, country_col in (
        ("快越达地址库", KYD_HEADERS, 0, 11),
        ("迈创地址库", MC_HEADERS, 1, 12),
    ):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for warehouse, country in (("YOW3", "CA"), ("LGB8", "US"), ("MQJ1", "US")):
            row = [None] * len(headers)
            row[code_col] = warehouse
            row[country_col] = country
            row[-1] = "K2J 7C7" if country == "CA" else 92551
            ws.append(row)
    del wb["Sheet"]
    wb.save(path)
    wb.close()
    monkeypatch.setattr(address_library, "CENTRAL_ADDRESS_PATH", path)
    monkeypatch.setattr(invoice_scan_service, "load_known_carriers", lambda: {"KYD", "MC"})
    return path


@pytest.mark.parametrize("carrier", ["KYD", "MC"])
def test_country_lookup_uses_each_carriers_address_sheet(country_address_book, carrier):
    countries = address_library.load_warehouse_countries(carrier)
    assert countries["YOW3"] == "CA"
    assert countries["LGB8"] == "US"
    assert countries["MQJ1"] == "US"


def test_kyd_preserves_canadian_postal_code(country_address_book):
    assert address_library.lookup_kyd_address("YOW3").zip_code == "K2J 7C7"


def test_source_country_and_legacy_country_subdirectory(tmp_path, monkeypatch):
    ini = tmp_path / "sources.ini"
    ini.write_text("[Source.KYD_CA]\nEnabled=TRUE\nCarrierCode=KYD\nCountryCode=ca\nPath=D:/KYD/未合并/加拿大发票\n", encoding="utf-8")
    monkeypatch.setattr(invoice_scan_service, "INVOICE_SOURCES_INI", ini)
    entry = invoice_scan_service.load_source_entries()[0]
    assert entry.country_code == "CA"
    assert invoice_scan_service.resolve_invoice_scan_directories(
        source_directory=Path("D:/KYD/未合并"), extra_directories=[],
        extra_only=False, country_code="CA",
    ) == [Path("D:/KYD/未合并/加拿大发票")]
    assert invoice_scan_service.resolve_invoice_scan_directories(
        source_directory=entry.path, extra_directories=[], extra_only=False,
        country_code="CA", source_country_code=entry.country_code,
    ) == [entry.path]


@pytest.mark.parametrize("warehouse,country", [("YOW3", "CA"), ("LGB8", "US"), ("MQJ1", "US")])
def test_scan_uses_warehouse_country_not_store_or_filename(tmp_path, country_address_book, warehouse, country):
    _write_meta_invoice(tmp_path / "misleading_US_name.xlsx", WarehouseCode=warehouse, StoreCode="美10")
    result = scan_original_invoices("20260819", [tmp_path], carrier_code="KYD", country_code=country)
    assert result.passed, result.errors
    assert result.country_code == country
    assert result.records[0].country_code == country


@pytest.mark.parametrize("warehouse", ["YOW3", "UNKNOWN"])
def test_scan_rejects_wrong_or_unknown_country(tmp_path, country_address_book, warehouse):
    _write_meta_invoice(tmp_path / "invoice.xlsx", WarehouseCode=warehouse)
    result = scan_original_invoices("20260819", [tmp_path], carrier_code="KYD", country_code="US")
    assert not result.passed
    assert any(warehouse in error for error in result.errors)


def test_scan_rejects_wrong_carrier_in_manual_directory(tmp_path, country_address_book):
    _write_meta_invoice(tmp_path / "invoice.xlsx", WarehouseCode="LGB8", CarrierCode="MC")
    result = scan_original_invoices("20260819", [tmp_path], carrier_code="KYD", country_code="US")
    assert not result.passed
    assert any("物流商" in error for error in result.errors)


def _country_record(country, carrier="KYD", **overrides):
    return make_record(country_code=country, carrier_code=carrier,
                       warehouse_code="YOW3" if country == "CA" else "LGB8", **overrides)


def test_batches_coexist_and_resume_by_carrier_and_country(tmp_path, monkeypatch):
    monkeypatch.setattr(batch_service, "BATCH_ROOT", tmp_path / "batches")
    us, _ = batch_service.ensure_batch_after_first_scan("20260819", make_result([_country_record("US")]))
    ca, _ = batch_service.ensure_batch_after_first_scan("20260819", make_result([_country_record("CA")]))
    mc, _ = batch_service.ensure_batch_after_first_scan("20260819", make_result([_country_record("US", "MC")]))
    for record, carrier, country in ((us, "KYD", "US"), (ca, "KYD", "CA"), (mc, "MC", "US")):
        assert batch_service.load_batch(record.directory).status == "FIRST_SCAN_PASSED"
        restored = batch_service.enforce_single_active("20260819", carrier_code=carrier, country_code=country)
        assert restored.batch_id == record.batch_id
        assert restored.snapshot["CountryCode"] == country


def test_forged_passed_scan_cannot_create_mixed_country_batch(tmp_path, monkeypatch):
    monkeypatch.setattr(batch_service, "BATCH_ROOT", tmp_path / "batches")
    with pytest.raises(batch_service.BatchError, match="国家"):
        batch_service.ensure_batch_after_first_scan("20260819", make_result([
            _country_record("US"), _country_record("CA", source_id="AAAAAAAA"),
        ]))
    assert not (tmp_path / "batches").exists()


def test_completing_country_keeps_other_country_and_carrier_outputs(tmp_path, monkeypatch):
    monkeypatch.setattr(batch_service, "BATCH_ROOT", tmp_path / "batches")
    monkeypatch.setattr("app.invoice_config.MERGE_RESULT_ROOT", tmp_path / "out")
    monkeypatch.setattr("app.services.invoice_fast_merge_service.QUICK_MERGE_ROOT", tmp_path / "quick")
    outputs = []
    for country, carrier in (("US", "KYD"), ("CA", "KYD"), ("US", "MC")):
        record, _ = batch_service.ensure_batch_after_first_scan("20260819", make_result([_country_record(country, carrier)]))
        output = batch_service.official_output_dir(record) / "invoice.xlsx"
        output.parent.mkdir(parents=True)
        output.write_bytes(b"invoice")
        outputs.append(output)
        batch_service.save_status(record, batch_service.STATUS_CHECK_PASSED)
        batch_service.complete_batch(record)
    assert all(path.exists() for path in outputs)


def test_merge_plan_rejects_mixed_legacy_snapshot_before_writing(tmp_path, monkeypatch, country_address_book):
    from app.services.merge_plan_service import build_merge_plan
    monkeypatch.setattr(batch_service, "BATCH_ROOT", tmp_path / "batches")
    batch, _ = batch_service.ensure_batch_after_first_scan("20260819", make_result([_country_record("US")]))
    batch.snapshot["Files"] = [vars(_country_record("US")), vars(_country_record("CA", source_id="AAAAAAAA"))]
    write_json(batch.directory / "quick_merge_manifest.json", {"Files": []})
    result = build_merge_plan(batch)
    assert not result.passed
    assert any("国家" in error for error in result.errors)
    assert not (batch.directory / "merge_plan.json").exists()


def test_content_merge_rejects_mixed_country_plan_before_output(tmp_path, monkeypatch):
    from app.services.invoice_content_merge_service import run_content_merge
    monkeypatch.setattr(batch_service, "BATCH_ROOT", tmp_path / "batches")
    batch, _ = batch_service.ensure_batch_after_first_scan("20260819", make_result([_country_record("US")]))
    write_json(batch.directory / "merge_plan.json", {"CountryCode": "US", "Groups": [
        {"carrier_code": "KYD", "country_code": "US", "warehouse_code": "LGB8"},
        {"carrier_code": "KYD", "country_code": "CA", "warehouse_code": "YOW3"},
    ]})
    result = run_content_merge(batch)
    assert not result.passed
    assert any("国家" in error for error in result.errors)
    assert result.output_files == []


def test_kyd_binary_output_keeps_canadian_postal_and_country(tmp_path, country_address_book):
    import xlrd
    from app.invoice_config import KYD_TEMPLATE_PATH
    from app.services.kyd_xls_io import copy_template, fill_kyd_template
    from tests.test_kyd_xls_fill import _one_line, _sample_photo_jpeg

    if not KYD_TEMPLATE_PATH.exists():
        pytest.skip("需要官方 KYD 模板")
    output = tmp_path / "KYD_YOW3_1箱.xls"
    copy_template(KYD_TEMPLATE_PATH, output)
    fill_kyd_template(output, channel="test", warehouse="YOW3", carton_count=1,
                      lines=[_one_line(1, _sample_photo_jpeg())])
    workbook = xlrd.open_workbook(str(output))
    try:
        sheet = workbook.sheet_by_name("模板")
        assert sheet.cell_value(11, 1) == "K2J 7C7"
        assert sheet.cell_value(12, 1) == "CA"
    finally:
        workbook.release_resources()


def test_changed_country_after_prepare_is_blocked_before_merge(tmp_path, monkeypatch, country_address_book):
    from openpyxl import load_workbook
    from app.services.invoice_prepare_service import run_prepare_merge
    from app.services.invoice_content_merge_service import run_content_merge
    monkeypatch.setattr(batch_service, "BATCH_ROOT", tmp_path / "batches")
    monkeypatch.setattr("app.services.invoice_fast_merge_service.QUICK_MERGE_ROOT", tmp_path / "quick")
    monkeypatch.setattr("app.services.merge_plan_service.MERGE_RESULT_ROOT", tmp_path / "out")
    monkeypatch.setattr("app.services.invoice_content_merge_service.LOCAL_TEMP_ROOT", tmp_path / "local")
    source_dir = tmp_path / "invoices"
    source_dir.mkdir()
    _write_meta_invoice(source_dir / "one.xlsx", WarehouseCode="LGB8")
    scan = scan_original_invoices("20260819", [source_dir], carrier_code="KYD", country_code="US")
    batch, _ = batch_service.ensure_batch_after_first_scan("20260819", scan)
    prepared = run_prepare_merge(batch)
    assert prepared.passed, prepared.errors
    plan = read_json(batch.directory / "merge_plan.json")
    assert plan["CountryCode"] == "US"
    assert plan["Groups"][0]["country_code"] == "US"
    copied = Path(plan["Groups"][0]["input_files"][0])
    wb = load_workbook(copied)
    for row in wb["_SystemMeta"].iter_rows():
        if row[0].value == "WarehouseCode":
            row[1].value = "YOW3"
    wb.save(copied)
    wb.close()
    result = run_content_merge(batch)
    assert not result.passed
    assert any("YOW3" in error and "国家" in error for error in result.errors)
    assert not list((tmp_path / "out").rglob("*.xls"))


def test_single_country_legacy_snapshot_can_be_restored(country_address_book):
    item = vars(make_record(warehouse_code="YOW3"))
    item.pop("country_code", None)
    restored = batch_service.scan_result_from_snapshot({"Files": [item], "BatchID": "legacy"})
    assert restored.country_code == "CA"
    assert restored.records[0].country_code == "CA"


def test_conflicting_address_country_is_not_guessed(tmp_path, country_address_book):
    from openpyxl import load_workbook
    wb = load_workbook(country_address_book)
    row = [None] * len(KYD_HEADERS)
    row[0], row[11] = "YOW3", "US"
    wb["快越达地址库"].append(row)
    wb.save(country_address_book)
    wb.close()
    _write_meta_invoice(tmp_path / "one.xlsx", WarehouseCode="YOW3")
    scan = scan_original_invoices("20260819", [tmp_path], country_code="CA", carrier_code="KYD")
    assert not scan.passed
    assert any("冲突" in error for error in scan.errors)


@pytest.mark.parametrize("old_warehouse,old_carrier", [("YOW3", "KYD"), ("LGB8", "MC")])
def test_old_generation_cannot_hide_wrong_country_or_carrier(tmp_path, country_address_book,
                                                            old_warehouse, old_carrier):
    _write_meta_invoice(tmp_path / "old.xlsx", WarehouseCode=old_warehouse,
                        CarrierCode=old_carrier, GeneratedAt="2026-08-19 10:00:00")
    _write_meta_invoice(tmp_path / "new.xlsx", WarehouseCode="LGB8",
                        CarrierCode="KYD", GeneratedAt="2026-08-19 11:00:00")
    scan = scan_original_invoices("20260819", [tmp_path], carrier_code="KYD", country_code="US")
    assert not scan.passed
    assert any("old.xlsx" in error for error in scan.errors)
