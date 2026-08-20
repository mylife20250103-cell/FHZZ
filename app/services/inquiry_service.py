from __future__ import annotations

import hashlib
import re
import shutil

from collections import defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

from app.inquiry_config import (
    INQUIRY_RESULT_ROOT,
    INQUIRY_TEMPLATE_PATH,
    INQUIRY_TEMPLATE_SHEET,
    INQUIRY_START_ROW,
    INQUIRY_WAREHOUSE_COLUMN,
    INQUIRY_WEIGHT_COLUMN,
    INQUIRY_CARTON_COLUMN,
    WAREHOUSE_PATTERN,
)
from app.invoice_config import (
    LOCAL_TEMP_ROOT,
    is_excel_junk_file,
    try_remove_excel_junk,
)


REQUIRED_COLUMNS = {
    "SchemaVersion",
    "DateID",
    "StoreCode",
    "PlanID",
    "SourceID",
    "WarehouseCode",
    "FBABatch",
    "CartonNumber",
    "WeightKG",
}


@dataclass
class InquirySummary:
    warehouse_code: str
    weight_kg: Decimal
    carton_count: int


@dataclass
class ScanResult:
    passed: bool
    file_count: int
    source_count: int
    carton_count: int
    warehouse_count: int
    total_weight: Decimal
    summaries: list[InquirySummary]
    errors: list[str]
    files: list[Path]


def excel_round_2(value) -> Decimal:
    """
    正数情况下等价于 Excel ROUND(value, 2)
    使用 ROUND_HALF_UP。
    """

    try:
        number = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        raise ValueError(
            f"不是有效数字：{value}"
        )

    return number.quantize(
        Decimal("0.01"),
        rounding=ROUND_HALF_UP,
    )


def sha256_file(path: Path) -> str:

    digest = hashlib.sha256()

    with path.open("rb") as f:

        for chunk in iter(
            lambda: f.read(1024 * 1024),
            b"",
        ):
            digest.update(chunk)

    return digest.hexdigest()


def collect_inquiry_files(
    directories: list[Path],
) -> list[Path]:

    results = []

    seen_paths = set()

    for directory in directories:

        if not directory.exists():
            continue

        if not directory.is_dir():
            continue

        for junk in directory.rglob("*"):
            if junk.is_file() and is_excel_junk_file(junk):
                try_remove_excel_junk(junk)

        for path in directory.rglob(
            "*_询价明细.xlsx"
        ):

            if is_excel_junk_file(path):
                try_remove_excel_junk(path)
                continue

            if path.name.startswith("~$"):
                continue

            resolved = path.resolve()

            if resolved in seen_paths:
                continue

            seen_paths.add(resolved)

            results.append(path)

    return sorted(
        results,
        key=lambda p: str(p).lower(),
    )


def read_inquiry_file(
    path: Path,
) -> tuple[list[dict], list[str]]:

    errors = []

    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
        )

    except Exception as exc:

        return [], [
            f"{path.name}：无法读取 Excel：{exc}"
        ]

    try:

        if "询价明细" not in workbook.sheetnames:

            return [], [
                f"{path.name}：缺少工作表【询价明细】"
            ]

        ws = workbook["询价明细"]

        header_row = next(
            ws.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True,
            )
        )

        headers = [
            str(value).strip()
            if value is not None
            else ""
            for value in header_row
        ]

        missing = (
            REQUIRED_COLUMNS
            - set(headers)
        )

        if missing:

            return [], [
                f"{path.name}：缺少必要字段："
                + ", ".join(sorted(missing))
            ]

        indexes = {
            name: headers.index(name)
            for name in REQUIRED_COLUMNS
        }

        rows = []

        for excel_row_number, values in enumerate(
            ws.iter_rows(
                min_row=2,
                values_only=True,
            ),
            start=2,
        ):

            # 全空行跳过
            if not any(
                value not in (None, "")
                for value in values
            ):
                continue

            row = {
                field: (
                    values[index]
                    if index < len(values)
                    else None
                )
                for field, index
                in indexes.items()
            }

            row["_ExcelRow"] = (
                excel_row_number
            )

            rows.append(row)

        if not rows:

            errors.append(
                f"{path.name}：询价明细没有数据"
            )

        return rows, errors

    finally:

        workbook.close()


def scan_inquiry_batch(
    date_id: str,
    directories: list[Path],
) -> ScanResult:

    errors = []

    files = collect_inquiry_files(
        directories
    )

    if not files:

        return ScanResult(
            passed=False,
            file_count=0,
            source_count=0,
            carton_count=0,
            warehouse_count=0,
            total_weight=Decimal("0.00"),
            summaries=[],
            errors=[
                "没有找到 *_询价明细.xlsx"
            ],
            files=[],
        )

    # ==========================================
    # 文件 SHA 防止同一份文件复制多份
    # ==========================================

    hash_to_file = {}

    for path in files:

        try:
            file_hash = sha256_file(path)

        except Exception as exc:

            errors.append(
                f"{path.name}："
                f"无法计算 SHA256：{exc}"
            )

            continue

        if file_hash in hash_to_file:

            errors.append(
                "发现内容完全相同的重复询价文件："
                f"{hash_to_file[file_hash]} "
                f"与 {path}"
            )

        else:
            hash_to_file[file_hash] = path

    # ==========================================
    # 全局数据
    # ==========================================

    source_to_file = {}

    carton_to_source = {}

    fba_to_warehouse = {}

    warehouse_weights = defaultdict(
        lambda: Decimal("0.00")
    )

    warehouse_cartons = defaultdict(set)

    global_cartons = set()

    source_ids = set()

    # ==========================================
    # 逐文件读取
    # ==========================================

    for path in files:

        rows, file_errors = (
            read_inquiry_file(path)
        )

        errors.extend(file_errors)

        if not rows:
            continue

        file_source_ids = set()
        file_store_codes = set()
        file_plan_ids = set()
        file_date_ids = set()

        for row in rows:

            excel_row = row["_ExcelRow"]

            def text(field):

                value = row.get(field)

                if value is None:
                    return ""

                return str(value).strip()

            schema_version = text(
                "SchemaVersion"
            )

            row_date = text("DateID")

            store_code = text(
                "StoreCode"
            )

            plan_id = text(
                "PlanID"
            )

            source_id = text(
                "SourceID"
            )

            warehouse = text(
                "WarehouseCode"
            )

            fba_batch = text(
                "FBABatch"
            )

            carton = text(
                "CartonNumber"
            )

            # ==================================
            # 基本字段
            # ==================================

            if schema_version != "1.0":

                errors.append(
                    f"{path.name} 第{excel_row}行："
                    f"SchemaVersion 必须为 1.0"
                )

            if row_date != date_id:

                errors.append(
                    f"{path.name} 第{excel_row}行："
                    f"DateID={row_date}，"
                    f"与当前选择日期 {date_id} 不一致"
                )

            if not source_id:

                errors.append(
                    f"{path.name} 第{excel_row}行："
                    "SourceID 为空"
                )

            if not store_code:

                errors.append(
                    f"{path.name} 第{excel_row}行："
                    "StoreCode 为空"
                )

            if not plan_id:

                errors.append(
                    f"{path.name} 第{excel_row}行："
                    "PlanID 为空"
                )

            if not warehouse:

                errors.append(
                    f"{path.name} 第{excel_row}行："
                    "WarehouseCode 为空"
                )

            elif not re.fullmatch(
                WAREHOUSE_PATTERN,
                warehouse,
            ):

                errors.append(
                    f"{path.name} 第{excel_row}行："
                    f"WarehouseCode 格式异常："
                    f"{warehouse}"
                )

            if not fba_batch:

                errors.append(
                    f"{path.name} 第{excel_row}行："
                    "FBABatch 为空"
                )

            if not carton:

                errors.append(
                    f"{path.name} 第{excel_row}行："
                    "CartonNumber 为空"
                )

            # ==================================
            # CartonNumber
            # ==================================

            if fba_batch and carton:

                expected_prefix = (
                    f"{fba_batch}U"
                )

                if not carton.startswith(
                    expected_prefix
                ):

                    errors.append(
                        f"{path.name} 第{excel_row}行："
                        f"CartonNumber 与 FBABatch "
                        f"不匹配：{carton}"
                    )

                else:

                    suffix = carton[
                        len(expected_prefix):
                    ]

                    if not re.fullmatch(
                        r"\d{6}",
                        suffix,
                    ):

                        errors.append(
                            f"{path.name} "
                            f"第{excel_row}行："
                            f"CartonNumber 后缀必须"
                            f"是6位数字：{carton}"
                        )

                    elif int(suffix) < 1:

                        errors.append(
                            f"{path.name} "
                            f"第{excel_row}行："
                            f"CartonNumber 不允许 "
                            f"U000000"
                        )

            # ==================================
            # 重量
            # ==================================

            try:

                weight = excel_round_2(
                    row.get("WeightKG")
                )

                if weight <= 0:

                    raise ValueError(
                        "必须大于0"
                    )

            except Exception:

                errors.append(
                    f"{path.name} 第{excel_row}行："
                    f"WeightKG 无效："
                    f"{row.get('WeightKG')}"
                )

                continue

            # ==================================
            # 文件内部一致性
            # ==================================

            file_source_ids.add(
                source_id
            )

            file_store_codes.add(
                store_code
            )

            file_plan_ids.add(
                plan_id
            )

            file_date_ids.add(
                row_date
            )

            # ==================================
            # 全局 Carton 唯一
            # ==================================

            if carton in global_cartons:

                previous_source = (
                    carton_to_source.get(
                        carton,
                        ""
                    )
                )

                errors.append(
                    f"CartonNumber 全局重复："
                    f"{carton} "
                    f"(SourceID="
                    f"{previous_source} / "
                    f"{source_id})"
                )

            else:

                global_cartons.add(
                    carton
                )

                carton_to_source[
                    carton
                ] = source_id

            # ==================================
            # 同 FBA 只能属于一个 Warehouse
            # ==================================

            if fba_batch:

                previous_warehouse = (
                    fba_to_warehouse.get(
                        fba_batch
                    )
                )

                if (
                    previous_warehouse
                    and
                    previous_warehouse
                    != warehouse
                ):

                    errors.append(
                        f"FBABatch {fba_batch} "
                        f"同时出现在仓库 "
                        f"{previous_warehouse} "
                        f"和 {warehouse}"
                    )

                else:

                    fba_to_warehouse[
                        fba_batch
                    ] = warehouse

            # ==================================
            # 汇总
            # ==================================

            warehouse_weights[
                warehouse
            ] += weight

            warehouse_cartons[
                warehouse
            ].add(carton)

        # ======================================
        # 单文件一致性
        # ======================================

        if len(file_source_ids) != 1:

            errors.append(
                f"{path.name}：同一文件中出现"
                "多个 SourceID"
            )

        if len(file_store_codes) != 1:

            errors.append(
                f"{path.name}：同一文件中出现"
                "多个 StoreCode"
            )

        if len(file_plan_ids) != 1:

            errors.append(
                f"{path.name}：同一文件中出现"
                "多个 PlanID"
            )

        if len(file_date_ids) != 1:

            errors.append(
                f"{path.name}：同一文件中出现"
                "多个 DateID"
            )

        if len(file_source_ids) == 1:

            source_id = next(
                iter(file_source_ids)
            )

            if source_id in source_to_file:

                errors.append(
                    f"同一 DateID 下 SourceID "
                    f"{source_id} 出现多个询价文件："
                    f"{source_to_file[source_id]} "
                    f"与 {path}"
                )

            else:

                source_to_file[
                    source_id
                ] = path

                source_ids.add(
                    source_id
                )

    # ==========================================
    # 生成预览
    # ==========================================

    summaries = []

    for warehouse in sorted(
        warehouse_weights.keys()
    ):

        summaries.append(
            InquirySummary(
                warehouse_code=warehouse,
                weight_kg=(
                    warehouse_weights[
                        warehouse
                    ].quantize(
                        Decimal("0.01")
                    )
                ),
                carton_count=len(
                    warehouse_cartons[
                        warehouse
                    ]
                ),
            )
        )

    total_weight = sum(
        (
            item.weight_kg
            for item in summaries
        ),
        Decimal("0.00"),
    )

    warehouse_carton_total = sum(
        item.carton_count
        for item in summaries
    )

    # ==========================================
    # 最终一致性检查
    # ==========================================

    if warehouse_carton_total != len(
        global_cartons
    ):

        errors.append(
            "仓库箱数之和与全局唯一 "
            "CartonNumber 数量不一致"
        )

    return ScanResult(
        passed=not errors,
        file_count=len(files),
        source_count=len(source_ids),
        carton_count=len(global_cartons),
        warehouse_count=len(summaries),
        total_weight=total_weight,
        summaries=summaries,
        errors=errors,
        files=files,
    )


def generate_inquiry_summary(
    date_id: str,
    summaries: list[InquirySummary],
) -> Path:
    """
    使用真实 Excel COM 生成询价汇总表。

    原则：
    1. 不修改中央模板原文件
    2. 先复制到临时文件
    3. Excel 自己维护 ListObject
    4. A:C 写仓库 / 重量 / 总箱数
    5. 保留 D:K 报价区域和公式
    6. 生成完成后重新打开验证
    7. 验证成功后才替换当天旧结果
    """

    import win32com.client as win32

    if not summaries:
        raise ValueError(
            "没有可生成的询价汇总数据"
        )

    if not INQUIRY_TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            f"询价模板不存在："
            f"{INQUIRY_TEMPLATE_PATH}"
        )

    # ==========================================
    # 输出目录
    # ==========================================

    output_directory = (
        INQUIRY_RESULT_ROOT / date_id
    )

    output_directory.mkdir(
        parents=True,
        exist_ok=True,
    )

    total_cartons = sum(
        item.carton_count
        for item in summaries
    )

    output_name = (
        f"{date_id}_询价汇总_"
        f"{total_cartons}箱.xlsx"
    )

    output_path = (
        output_directory / output_name
    )

    LOCAL_TEMP_ROOT.mkdir(parents=True, exist_ok=True)
    temp_path = LOCAL_TEMP_ROOT / f"inquiry_tmp_{date_id}.xlsx"

    # 防止上次异常遗留
    if temp_path.exists():
        try:
            temp_path.unlink()
        except OSError as exc:
            raise PermissionError(
                f"无法清理临时文件："
                f"{temp_path}\n\n{exc}"
            )

    # ==========================================
    # 复制中央模板
    # ==========================================

    shutil.copy2(
        INQUIRY_TEMPLATE_PATH,
        temp_path,
    )

    excel = None
    workbook = None

    try:

        # ======================================
        # 独立 Excel 实例
        # ======================================

        excel = win32.DispatchEx(
            "Excel.Application"
        )

        excel.Visible = False
        excel.DisplayAlerts = False
        excel.ScreenUpdating = False
        excel.EnableEvents = False

        workbook = excel.Workbooks.Open(
            str(temp_path),
            UpdateLinks=0,
            ReadOnly=False,
        )

        try:

            ws = workbook.Worksheets(
                INQUIRY_TEMPLATE_SHEET
            )

        except Exception:

            raise ValueError(
                "询价模板缺少工作表："
                f"{INQUIRY_TEMPLATE_SHEET}"
            )

        # ======================================
        # 获取两个 Excel Table
        # ======================================

        try:
            source_table = (
                ws.ListObjects("第三步_1")
            )

        except Exception:

            raise ValueError(
                "询价模板缺少 Excel 表格："
                "第三步_1"
            )

        try:
            quote_table = (
                ws.ListObjects("第三步")
            )

        except Exception:

            raise ValueError(
                "询价模板缺少 Excel 表格："
                "第三步"
            )

        # ======================================
        # 校验表头
        # ======================================

        source_headers = [
            source_table.ListColumns(i).Name
            for i in range(
                1,
                source_table.ListColumns.Count + 1
            )
        ]

        expected_source_headers = [
            "仓库",
            "重量",
            "总箱数",
        ]

        if (
            source_headers
            != expected_source_headers
        ):

            raise ValueError(
                "询价模板【第三步_1】字段异常。\n\n"
                f"期望：{expected_source_headers}\n"
                f"实际：{source_headers}"
            )

        if quote_table.ListColumns.Count != 8:

            raise ValueError(
                "询价模板【第三步】"
                "应为8列报价区域"
            )

        # ======================================
        # 数据行容量
        # ======================================

        current_source_rows = (
            source_table.ListRows.Count
        )

        current_quote_rows = (
            quote_table.ListRows.Count
        )

        required_rows = len(
            summaries
        )

        # 两个 Table 保持相同行数。
        #
        # 当前模板：
        # 第三步_1 = 17行
        # 第三步   = 28行
        #
        # 因此正常会统一成28行。
        desired_rows = max(
            required_rows,
            current_source_rows,
            current_quote_rows,
        )

        if desired_rows < 1:
            desired_rows = 1

        last_row = (
            desired_rows + 1
        )

        # ======================================
        # 调整两个 Table
        # ======================================

        source_table.Resize(
            ws.Range(
                f"A1:C{last_row}"
            )
        )

        quote_table.Resize(
            ws.Range(
                f"D1:K{last_row}"
            )
        )

        # ======================================
        # 清空旧 A:C 数据
        # ======================================

        ws.Range(
            f"A2:C{last_row}"
        ).ClearContents()

        # ======================================
        # 写入仓库汇总
        # ======================================

        for index, summary in enumerate(
            summaries,
            start=2,
        ):

            ws.Cells(
                index,
                1,
            ).Value = (
                summary.warehouse_code
            )

            ws.Cells(
                index,
                2,
            ).Value = float(
                summary.weight_kg
            )

            ws.Cells(
                index,
                3,
            ).Value = (
                summary.carton_count
            )

        # ======================================
        # 重量格式
        # ======================================

        ws.Range(
            f"B2:B{last_row}"
        ).NumberFormat = "0.00"

        # ======================================
        # 保证 G / K 公式继续存在
        # ======================================

        if desired_rows >= 1:

            g2_formula = (
                ws.Range("G2").Formula
            )

            k2_formula = (
                ws.Range("K2").Formula
            )

            if not g2_formula:
                raise ValueError(
                    "模板 G2 总价公式为空"
                )

            if not k2_formula:
                raise ValueError(
                    "模板 K2 总价公式为空"
                )

            if desired_rows > 1:

                ws.Range(
                    f"G2:G{last_row}"
                ).FillDown()

                ws.Range(
                    f"K2:K{last_row}"
                ).FillDown()

        # ======================================
        # 保存
        # ======================================

        workbook.Save()

        workbook.Close(
            SaveChanges=True
        )

        workbook = None

    except Exception:

        if workbook is not None:

            try:
                workbook.Close(
                    SaveChanges=False
                )
            except Exception:
                pass

        raise

    finally:

        if excel is not None:

            try:
                excel.DisplayAlerts = False
                excel.Quit()
            except Exception:
                pass

            excel = None

    # ==========================================
    # 第二次：COM 回读验证
    # ==========================================

    check_excel = None
    check_book = None

    try:

        check_excel = win32.DispatchEx(
            "Excel.Application"
        )

        check_excel.Visible = False
        check_excel.DisplayAlerts = False

        check_book = (
            check_excel.Workbooks.Open(
                str(temp_path),
                UpdateLinks=0,
                ReadOnly=True,
            )
        )

        check_ws = (
            check_book.Worksheets(
                INQUIRY_TEMPLATE_SHEET
            )
        )

        check_source_table = (
            check_ws.ListObjects(
                "第三步_1"
            )
        )

        check_quote_table = (
            check_ws.ListObjects(
                "第三步"
            )
        )

        # ======================================
        # Table 行数一致
        # ======================================

        if (
            check_source_table.ListRows.Count
            !=
            check_quote_table.ListRows.Count
        ):

            raise ValueError(
                "生成后校验失败："
                "两个 Excel Table 行数不一致"
            )

        # ======================================
        # 逐仓库验证
        # ======================================

        for row_number, summary in enumerate(
            summaries,
            start=2,
        ):

            warehouse = (
                check_ws.Cells(
                    row_number,
                    1,
                ).Value
            )

            weight = (
                check_ws.Cells(
                    row_number,
                    2,
                ).Value
            )

            cartons = (
                check_ws.Cells(
                    row_number,
                    3,
                ).Value
            )

            if (
                str(warehouse).strip()
                !=
                summary.warehouse_code
            ):

                raise ValueError(
                    "生成后仓库校验失败："
                    f"{summary.warehouse_code}"
                )

            actual_weight = (
                excel_round_2(weight)
            )

            if (
                actual_weight
                !=
                summary.weight_kg
            ):

                raise ValueError(
                    "生成后重量校验失败："
                    f"{summary.warehouse_code}"
                )

            if int(cartons) != (
                summary.carton_count
            ):

                raise ValueError(
                    "生成后箱数校验失败："
                    f"{summary.warehouse_code}"
                )

        check_book.Close(
            SaveChanges=False
        )

        check_book = None

    finally:

        if check_book is not None:

            try:
                check_book.Close(
                    SaveChanges=False
                )
            except Exception:
                pass

        if check_excel is not None:

            try:
                check_excel.Quit()
            except Exception:
                pass

    # ==========================================
    # 正式替换
    #
    # 注意：
    # 不先删除旧结果。
    #
    # temp 在本地生成并通过验证后，再复制到中央结果目录，
    # 避免 Excel 在 OneDrive 上留下 .~tmp_ 文件。
    # ==========================================

    try:

        shutil.copy2(temp_path, output_path)
        try:
            temp_path.unlink()
        except OSError:
            pass

    except PermissionError:

        raise PermissionError(
            "无法替换询价结果文件。\n\n"
            "可能原因：当天的询价汇总表"
            "目前正在 Excel 中打开。\n\n"
            "请关闭该文件后重新生成。\n\n"
            f"{output_path}"
        )

    for folder in (output_directory, LOCAL_TEMP_ROOT):
        for junk in folder.glob(".~tmp*"):
            try_remove_excel_junk(junk)

    # ==========================================
    # 删除同日其他旧结果
    # ==========================================

    for old_file in (
        output_directory.glob(
            f"{date_id}_询价汇总_*箱.xlsx"
        )
    ):

        if old_file == output_path:
            continue

        try:
            old_file.unlink()

        except OSError:
            # 不影响本次正式结果
            pass

    return output_path