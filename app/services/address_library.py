from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from app.invoice_config import CENTRAL_ADDRESS_PATH


KYD_SHEET = "快越达地址库"
MC_SHEET = "迈创地址库"

KYD_HEADER = "地址编码*"
MC_CODE_HEADER = "地址编码*"
MC_COLUMNS = 14


@dataclass
class AddressRow:
    warehouse: str
    contact: str
    company: str
    phone: float | None
    address1: str
    address2: str
    city: str
    state: str
    country: str
    zip_code: float | None


def _error(message: str):
    from app.services.invoice_adapters.base import AdapterError

    return AdapterError(message)


def _as_float(value) -> float | None:
    if value in ("", None):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _as_text(value) -> str:
    if value in ("", None):
        return ""
    return str(value).strip()


def _open_central(path: Path | None = None):
    source = path or CENTRAL_ADDRESS_PATH
    if not source.exists():
        raise _error(f"中央地址库不存在：{source}")
    try:
        return load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise _error(f"无法读取中央地址库：{exc}") from exc


def lookup_kyd_address(
    warehouse: str,
    path: Path | None = None,
) -> AddressRow:
    """从中央【快越达地址库】按地址编码取仓，供快越达只写「模板」表。"""

    target = (warehouse or "").strip().upper()
    if not target:
        raise _error("快越达合并没有仓库编码")

    workbook = _open_central(path)
    try:
        if KYD_SHEET not in workbook.sheetnames:
            raise _error(f"中央地址库缺少工作表【{KYD_SHEET}】")
        sheet = workbook[KYD_SHEET]
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if not row:
                continue
            code = _as_text(row[0] if len(row) > 0 else None)
            if index == 0:
                if KYD_HEADER not in code:
                    raise _error(
                        f"【{KYD_SHEET}】首列应为 {KYD_HEADER}，实际是 {code or '空'}"
                    )
                continue
            if not code:
                break
            if code.upper() != target:
                continue
            return AddressRow(
                warehouse=code,
                contact=_as_text(row[2] if len(row) > 2 else None),
                company=_as_text(row[3] if len(row) > 3 else None),
                phone=_as_float(row[4] if len(row) > 4 else None),
                address1=_as_text(row[6] if len(row) > 6 else None),
                address2=_as_text(row[7] if len(row) > 7 else None),
                city=_as_text(row[9] if len(row) > 9 else None),
                state=_as_text(row[10] if len(row) > 10 else None),
                country=_as_text(row[11] if len(row) > 11 else None),
                zip_code=_as_float(row[12] if len(row) > 12 else None),
            )
    finally:
        workbook.close()

    raise _error(f"中央地址库没有仓库 {warehouse}")


def load_mc_address_rows(path: Path | None = None) -> list[list[object]]:
    """读出中央【迈创地址库】整表（含表头），写入合并复制件的「地址库」。"""

    workbook = _open_central(path)
    try:
        if MC_SHEET not in workbook.sheetnames:
            raise _error(f"中央地址库缺少工作表【{MC_SHEET}】")
        sheet = workbook[MC_SHEET]
        rows: list[list[object]] = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            values = list(row[:MC_COLUMNS]) if row else []
            while len(values) < MC_COLUMNS:
                values.append(None)
            if index == 0:
                header = _as_text(values[1] if len(values) > 1 else None)
                if MC_CODE_HEADER not in header and MC_CODE_HEADER not in _as_text(
                    values[0]
                ):
                    raise _error(
                        f"【{MC_SHEET}】缺少列 {MC_CODE_HEADER}"
                    )
                rows.append(values)
                continue
            code = _as_text(values[1])
            if not code:
                break
            rows.append(values)
        if len(rows) < 2:
            raise _error(f"【{MC_SHEET}】没有地址数据")
        return rows
    finally:
        workbook.close()
