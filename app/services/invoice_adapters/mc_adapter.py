from __future__ import annotations

import atexit
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook

from app.excel_com import close_workbook, open_workbook
from app.invoice_config import CENTRAL_ADDRESS_PATH, MC_TEMPLATE_PATH, appdata_root
from app.services.address_library import load_mc_address_rows
from app.services.invoice_adapters.base import (
    AdapterError,
    AdapterTemplateConfig,
    InvoiceMergeAdapter,
    ProductFieldSpec,
)


_LOOKUP_ROW = re.compile(r"(\$?N)\$?\d+", re.IGNORECASE)
_mc_excel = None


def close_mc_excel() -> None:
    """一批迈创合并结束后关闭共用的 Excel。没启动过则什么都不做。"""

    global _mc_excel
    excel = _mc_excel
    _mc_excel = None
    if excel is None:
        return
    try:
        excel.DisplayAlerts = False
        excel.Quit()
    except Exception:
        pass


atexit.register(close_mc_excel)


def _mc_calculator():

    global _mc_excel
    if _mc_excel is None:
        import win32com.client as win32

        excel = win32.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AskToUpdateLinks = False
        excel.AlertBeforeOverwriting = False
        excel.ScreenUpdating = False
        excel.EnableEvents = False
        _mc_excel = excel
    return _mc_excel


def _calculate_local(path: Path) -> None:

    excel = _mc_calculator()
    workbook = open_workbook(excel, path, read_only=False)
    try:
        excel.Calculate()
        workbook.Save()
    finally:
        close_workbook(workbook, save=True)


class McAdapter(InvoiceMergeAdapter):
    """
    迈创适配器。

    源发票工作表【迈创发票】，表头第17行，明细从第18行起。
    产品图是 R 列 URL，不转浮动图。
    内容合并复制中央【迈创发票模板.xlsx】，只填复制件的【模板】表。
    地址库在本机缓存里注入一次，各仓用 openpyxl 写明细。
    整批共用一个 Excel，只负责计算公式。官方模板文件本身不改。
    """

    carrier_code = "MC"
    template_version = "1.0"
    carrier_name = "迈创"

    SOURCE_SHEET = "迈创发票"
    OUTPUT_SHEET = "模板"
    LAST_DATA_COL = 22
    EMPTY_CARTON_STOP = 8
    FLAG_CELLS = ("F1", "F2", "F3", "F4", "F5", "F6", "F8")
    ADDRESS_SHEET = "地址库"
    ADDRESS_LOOKUP_CELLS = ("B6", "B9", "B10", "B11", "B12")
    ADDRESS_COLUMNS = 14

    def template_config(self) -> AdapterTemplateConfig:

        return AdapterTemplateConfig(
            sheet_name=self.SOURCE_SHEET,
            header_row=17,
            data_start_row=18,
            channel_cell="B2",
            warehouse_cell="B3",
            carton_count_cell="B16",
            product_fields=[
                ProductFieldSpec(
                    "ProductEN",
                    "产品英文品名",
                    7,
                    "产品英文品名*",
                ),
                ProductFieldSpec(
                    "ProductCN",
                    "产品中文品名",
                    8,
                    "产品中文品名*",
                ),
                ProductFieldSpec(
                    "MaterialEN",
                    "产品材质",
                    11,
                    "产品材质*",
                ),
                ProductFieldSpec(
                    "MaterialCN",
                    "产品材质",
                    11,
                    "产品材质*",
                ),
            ],
            notes=(
                "MC 从中央【迈创发票模板.xlsx】复制后填写【模板】表。"
                "B2 渠道用模板自带名称，不覆盖源发票旧服务名。"
                "「地址库」注入本机缓存的复制件，不改官方模板。"
                "产品图片在 R 列链接，不是浮动图。"
            ),
        )

    def output_extension(self) -> str:

        return ".xlsx"

    def validate_source(self, workbook_path, meta: dict) -> list[str]:

        errors = []
        path = Path(workbook_path)
        cfg = self.template_config()

        if meta.get("CarrierCode", "").upper() != self.carrier_code:
            errors.append(f"{path.name}：CarrierCode 不是 MC")

        if meta.get("TemplateVersion") != self.template_version:
            errors.append(
                f"{path.name}：TemplateVersion 不是 {self.template_version}"
            )

        try:
            workbook = load_workbook(
                path,
                read_only=True,
                data_only=False,
            )
        except Exception as exc:
            errors.append(f"{path.name}：无法打开：{exc}")
            return errors

        try:
            if cfg.sheet_name not in workbook.sheetnames:
                errors.append(
                    f"{path.name}：缺少工作表【{cfg.sheet_name}】"
                )
        finally:
            workbook.close()

        return errors

    def diagnose_source(self, workbook_path) -> list[str]:

        path = Path(workbook_path)
        cfg = self.template_config()
        try:
            workbook = load_workbook(path, data_only=False)
        except Exception as exc:
            return [f"{path.name}：无法打开：{exc}"]

        issues: list[str] = []
        try:
            if self.SOURCE_SHEET not in workbook.sheetnames:
                return [f"{path.name}：缺少工作表【{self.SOURCE_SHEET}】"]
            ws = workbook[self.SOURCE_SHEET]
            cartons = self._count_data_rows_openpyxl(ws, cfg.data_start_row)
            if not cartons:
                issues.append(f"{path.name}：没有找到明细行")
            warehouse = ws[cfg.warehouse_cell].value
            if warehouse in (None, "") or not str(warehouse).strip():
                issues.append(f"{path.name}：{cfg.warehouse_cell} 仓库为空")
            empty_run = 0
            last = min(ws.max_row or cfg.data_start_row, cfg.data_start_row + 200)
            for row in range(cfg.data_start_row, last + 1):
                carton = ws.cell(row, 1).value
                if carton in (None, ""):
                    empty_run += 1
                    if empty_run >= self.EMPTY_CARTON_STOP:
                        break
                    continue
                empty_run = 0
                label = f"{path.name} 第{row}行 {str(carton).strip()}"
                if ws.cell(row, 7).value in (None, ""):
                    issues.append(f"{label}：产品英文品名为空")
                if ws.cell(row, 8).value in (None, ""):
                    issues.append(f"{label}：产品中文品名为空")
        except Exception as exc:
            issues.append(f"{path.name}：自检异常：{exc}")
        finally:
            workbook.close()

        return issues

    def _is_data_carton(self, value) -> bool:

        return value not in (None, "")

    def _count_data_rows_openpyxl(self, worksheet, start_row: int) -> list[str]:

        cartons = []
        empty_run = 0
        last = min(worksheet.max_row or start_row, start_row + 200)

        for row in range(start_row, last + 1):
            carton = worksheet.cell(row, 1).value

            if self._is_data_carton(carton):
                cartons.append(str(carton).strip())
                empty_run = 0
            else:
                empty_run += 1
                if empty_run >= self.EMPTY_CARTON_STOP:
                    break

        return cartons

    def _prepared_template(self) -> Path:

        if not MC_TEMPLATE_PATH.exists():
            raise AdapterError(f"迈创发票模板不存在：{MC_TEMPLATE_PATH}")

        cache = appdata_root() / "cache"
        cache.mkdir(parents=True, exist_ok=True)
        dest = cache / "mc_address_template.xlsx"
        stamp_path = cache / "mc_address_template.stamp"
        template_stat = MC_TEMPLATE_PATH.stat()
        address_stat = CENTRAL_ADDRESS_PATH.stat()
        stamp = (
            f"{template_stat.st_mtime_ns}:{template_stat.st_size}:"
            f"{address_stat.st_mtime_ns}:{address_stat.st_size}"
        )
        if dest.exists() and stamp_path.exists():
            if stamp_path.read_text(encoding="utf-8") == stamp:
                return dest

        rows = load_mc_address_rows()
        shutil.copy2(MC_TEMPLATE_PATH, dest)
        workbook = load_workbook(dest)
        try:
            if self.ADDRESS_SHEET not in workbook.sheetnames:
                raise AdapterError(f"迈创发票模板缺少工作表【{self.ADDRESS_SHEET}】")
            if self.OUTPUT_SHEET not in workbook.sheetnames:
                raise AdapterError(f"迈创发票模板缺少工作表【{self.OUTPUT_SHEET}】")
            self._write_address_sheet(workbook[self.ADDRESS_SHEET], rows)
            self._retarget_address_lookups(workbook[self.OUTPUT_SHEET], len(rows))
            workbook.save(dest)
        finally:
            workbook.close()
        stamp_path.write_text(stamp, encoding="utf-8")
        return dest

    def _write_address_sheet(self, worksheet, rows: list[list[object]]) -> None:

        if not rows:
            raise AdapterError("迈创中央地址库是空的")
        old_rows = worksheet.max_row or 1
        clear_to = max(len(rows), 990)
        if clear_to < old_rows <= 3000:
            clear_to = old_rows
        for row_index in range(1, clear_to + 1):
            for col in range(1, self.ADDRESS_COLUMNS + 1):
                worksheet.cell(row_index, col).value = None
        for row_index, values in enumerate(rows, start=1):
            for col in range(1, self.ADDRESS_COLUMNS + 1):
                value = values[col - 1] if col - 1 < len(values) else None
                worksheet.cell(row_index, col).value = value

    def _retarget_address_lookups(self, worksheet, last_row: int) -> None:

        last_row = max(int(last_row), 2)
        for coord in self.ADDRESS_LOOKUP_CELLS:
            formula = worksheet[coord].value
            if not isinstance(formula, str) or "地址库" not in formula:
                continue
            worksheet[coord].value = _LOOKUP_ROW.sub(
                rf"\g<1>${last_row}",
                formula,
                count=1,
            )

    def _read_source(self, path: Path) -> tuple[dict[str, object], list[list]]:

        workbook = load_workbook(path, data_only=False)
        try:
            if self.SOURCE_SHEET not in workbook.sheetnames:
                raise AdapterError(f"{path.name}：缺少工作表【{self.SOURCE_SHEET}】")
            sheet = workbook[self.SOURCE_SHEET]
            flags = {coord: sheet[coord].value for coord in self.FLAG_CELLS}
            start = self.template_config().data_start_row
            rows: list[list] = []
            empty_run = 0
            for row_index in range(start, start + 200):
                values = [
                    sheet.cell(row_index, col).value
                    for col in range(1, self.LAST_DATA_COL + 1)
                ]
                if self._is_data_carton(values[0]):
                    rows.append(values)
                    empty_run = 0
                else:
                    empty_run += 1
                    if empty_run >= self.EMPTY_CARTON_STOP:
                        break
            return flags, rows
        finally:
            workbook.close()

    def merge_group(self, group: dict, output_path) -> None:

        inputs = [Path(item) for item in group["input_files"]]
        if not inputs:
            raise AdapterError("MC 合并分组没有输入文件")

        output_path = Path(output_path)
        if output_path.suffix.lower() != ".xlsx":
            output_path = output_path.with_suffix(".xlsx")

        template_mtime = MC_TEMPLATE_PATH.stat().st_mtime
        prepared = self._prepared_template()
        work = appdata_root() / "cache"
        work.mkdir(parents=True, exist_ok=True)
        local = work / f"mc_fill_{output_path.name}"
        shutil.copy2(prepared, local)

        flags: dict[str, object] = {}
        data_rows: list[list] = []
        for index, source in enumerate(inputs):
            source_flags, rows = self._read_source(source)
            if not rows:
                raise AdapterError(f"{source.name}：没有找到明细行")
            if index == 0:
                flags = source_flags
            data_rows.extend(rows)

        cfg = self.template_config()
        workbook = load_workbook(local)
        try:
            sheet = workbook[self.OUTPUT_SHEET]
            for coord, value in flags.items():
                if value not in (None, ""):
                    sheet[coord].value = value
            for offset, values in enumerate(data_rows):
                dest_row = cfg.data_start_row + offset
                for col, value in enumerate(values, start=1):
                    sheet.cell(dest_row, col).value = value
            sheet[cfg.warehouse_cell].value = str(group["warehouse_code"]).strip()
            sheet[cfg.carton_count_cell].value = len(group["carton_numbers"])
            workbook.save(local)
        finally:
            workbook.close()

        _calculate_local(local)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(local, output_path)

        if MC_TEMPLATE_PATH.stat().st_mtime != template_mtime:
            raise AdapterError("官方迈创发票模板被改到了，请检查")

    def validate_output(self, output_path, group: dict) -> list[str]:

        errors = []
        path = Path(output_path)
        cfg = self.template_config()
        warehouse = str(group["warehouse_code"]).strip()

        if not path.exists():
            return [f"输出文件不存在：{path}"]
        if path.suffix.lower() != ".xlsx":
            return [f"{path.name}：MC 输出必须是 .xlsx"]
        if path.name.startswith("~$"):
            return [f"输出是临时锁文件：{path.name}"]

        try:
            workbook = load_workbook(path, data_only=False)
        except Exception as exc:
            return [f"{path.name}：无法打开：{exc}"]

        try:
            names = workbook.sheetnames
            for required in (self.OUTPUT_SHEET, "渠道列表", "地址库"):
                if required not in names:
                    errors.append(f"{path.name}：缺少工作表【{required}】")
            if errors:
                return errors

            if "_SystemMeta" in names:
                errors.append(f"{path.name}：不应保留 _SystemMeta")
            if "_MergeMeta" in names:
                errors.append(f"{path.name}：不应写入 _MergeMeta")

            ws = workbook[self.OUTPUT_SHEET]
            channel = ws[cfg.channel_cell].value
            if channel in (None, ""):
                errors.append(f"{path.name}：B2 渠道不能为空")

            b4 = ws["B4"].value
            if not (isinstance(b4, str) and b4.startswith("=")):
                errors.append(f"{path.name}：B4 必须保留公式")

            actual_wh = "" if ws[cfg.warehouse_cell].value is None else str(
                ws[cfg.warehouse_cell].value
            ).strip()
            if actual_wh.upper() != warehouse.upper():
                errors.append(
                    f"{path.name}：仓库单元格不是 {warehouse}"
                )

            carton_cell = ws[cfg.carton_count_cell].value
            try:
                carton_value = int(float(carton_cell))
            except (TypeError, ValueError):
                carton_value = None
            if carton_value != len(group["carton_numbers"]):
                errors.append(
                    f"{path.name}：{cfg.carton_count_cell} "
                    f"应为 {len(group['carton_numbers'])}"
                )

            found_cartons = self._count_data_rows_openpyxl(
                ws,
                cfg.data_start_row,
            )
            expected_cartons = [
                str(item).strip()
                for item in group["carton_numbers"]
            ]
            if len(found_cartons) != len(expected_cartons):
                errors.append(
                    f"{path.name}：明细行 {len(found_cartons)}，"
                    f"应为 {len(expected_cartons)} 箱"
                )
            elif sorted(found_cartons) != sorted(expected_cartons):
                errors.append(
                    f"{path.name}：明细箱号与 MergePlan 不一致"
                )

            for offset, carton in enumerate(found_cartons):
                url = ws.cell(cfg.data_start_row + offset, 18).value
                if url in (None, ""):
                    errors.append(
                        f"{path.name}：{carton} 的 R 列图片链接为空"
                    )

        finally:
            workbook.close()

        return errors
