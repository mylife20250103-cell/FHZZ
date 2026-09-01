from __future__ import annotations

import struct
from pathlib import Path
from urllib.error import URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook

from app.invoice_config import KYD_TEMPLATE_PATH
from app.services.invoice_adapters.base import (
    AdapterError,
    AdapterTemplateConfig,
    InvoiceMergeAdapter,
    ProductFieldSpec,
)
from app.services.kyd_xls_io import (
    KydLine,
    copy_template,
    fill_kyd_template,
    iter_biff,
    read_workbook_stream,
)
from app.services.address_library import lookup_kyd_address


class KydAdapter(InvoiceMergeAdapter):

    carrier_code = "KYD"
    template_version = "1.0"
    carrier_name = "快越达"

    SOURCE_SHEET = "快越达发票"
    OUTPUT_SHEET = "模板"
    LAST_DATA_COL = 20

    def template_config(self) -> AdapterTemplateConfig:

        return AdapterTemplateConfig(
            sheet_name=self.SOURCE_SHEET,
            header_row=29,
            data_start_row=30,
            channel_cell="B4",
            warehouse_cell="B5",
            carton_count_cell="B26",
            product_fields=[
                ProductFieldSpec(
                    "ProductEN",
                    "产品英文品名",
                    3,
                    "产品英文品名*",
                ),
                ProductFieldSpec(
                    "ProductCN",
                    "产品中文品名",
                    4,
                    "产品中文品名*",
                ),
                ProductFieldSpec(
                    "MaterialEN",
                    "产品材质",
                    13,
                    "产品材质*",
                ),
                ProductFieldSpec(
                    "MaterialCN",
                    "产品材质",
                    13,
                    "产品材质*",
                ),
            ],
            notes=(
                "KYD 从中央【快越达发票模版.xls】复制后填写【模板】表。"
                "地址从【各物流地址库.xlsx】注入 B6–B13，不改模板里的「地址库」表。"
                "源发票 S 列是图片网址，输出为浮动图。"
            ),
        )

    def output_extension(self) -> str:

        return ".xls"

    def validate_source(self, workbook_path, meta: dict) -> list[str]:

        errors = []
        path = Path(workbook_path)
        cfg = self.template_config()

        if meta.get("CarrierCode", "").upper() != self.carrier_code:
            errors.append(f"{path.name}：CarrierCode 不是 KYD")

        if meta.get("TemplateVersion") != self.template_version:
            errors.append(
                f"{path.name}：TemplateVersion 不是 {self.template_version}"
            )

        try:
            workbook = load_workbook(
                path,
                read_only=True,
                data_only=False,
            )
        except Exception as exc:
            return [f"{path.name}：无法打开：{exc}"]

        try:
            if cfg.sheet_name not in workbook.sheetnames:
                errors.append(
                    f"{path.name}：缺少工作表【{cfg.sheet_name}】"
                )
        finally:
            workbook.close()

        return errors

    def diagnose_source(self, workbook_path) -> list[str]:

        path = Path(workbook_path)
        issues: list[str] = []
        try:
            workbook = load_workbook(path, data_only=False)
        except Exception as exc:
            return [f"{path.name}：无法打开：{exc}"]

        try:
            if self.SOURCE_SHEET not in workbook.sheetnames:
                return [f"{path.name}：缺少工作表【{self.SOURCE_SHEET}】"]
            ws = workbook[self.SOURCE_SHEET]
            if _blank(ws["B4"].value):
                issues.append(f"{path.name}：B4 渠道为空")
            if _blank(ws["B5"].value):
                issues.append(f"{path.name}：B5 仓库为空")
            rows = 0
            for row in range(30, (ws.max_row or 30) + 1):
                carton = ws.cell(row, 1).value
                product = ws.cell(row, 3).value
                if _blank(carton) and _blank(product):
                    continue
                rows += 1
                carton_text = "" if carton is None else str(carton).strip()
                label = f"{path.name} 第{row}行"
                if carton_text:
                    label += f" {carton_text}"
                if _blank(carton):
                    issues.append(f"{label}：箱号为空")
                if _blank(product):
                    issues.append(f"{label}：产品英文品名为空")
                if _blank(ws.cell(row, 4).value):
                    issues.append(f"{label}：产品中文品名为空")
                if not _is_number(ws.cell(row, 6).value):
                    issues.append(f"{label}：申报单价不是数字")
                if not _is_number(ws.cell(row, 7).value):
                    issues.append(f"{label}：申报数量不是数字")
                if _blank(ws.cell(row, 10).value):
                    issues.append(f"{label}：海关编码为空")
                url = "" if ws.cell(row, 19).value is None else str(
                    ws.cell(row, 19).value
                ).strip()
                if not url.lower().startswith(("http://", "https://")):
                    issues.append(f"{label}：产品图片网址为空或不是网址")
                else:
                    try:
                        download_product_image(url, path.name)
                    except AdapterError as exc:
                        text = str(exc)
                        prefix = f"{path.name}："
                        if text.startswith(prefix):
                            text = text[len(prefix):]
                        issues.append(f"{label}：{text}")
            if rows == 0:
                issues.append(f"{path.name}：没有找到明细行")
        except Exception as exc:
            issues.append(f"{path.name}：自检异常：{exc}")
        finally:
            workbook.close()

        return issues

    def _source_rows(self, path: Path) -> tuple[str, list[list[object]], list[str]]:

        workbook = load_workbook(path, data_only=False)
        try:
            if self.SOURCE_SHEET not in workbook.sheetnames:
                raise AdapterError(f"{path.name}：缺少工作表【{self.SOURCE_SHEET}】")
            ws = workbook[self.SOURCE_SHEET]
            channel = ws["B4"].value
            rows = []
            urls = []
            for row in range(30, (ws.max_row or 30) + 1):
                carton = ws.cell(row, 1).value
                product = ws.cell(row, 3).value
                if carton in (None, "") and product in (None, ""):
                    continue
                values = [ws.cell(row, col).value for col in range(1, 21)]
                formula = values[8]
                unit = values[5]
                qty = values[6]
                if _is_number(unit) and _is_number(qty):
                    values[8] = float(unit) * float(qty)
                elif isinstance(formula, str) and formula.startswith("="):
                    values[8] = None
                values[18] = None
                rows.append(values)
                url = ws.cell(row, 19).value
                urls.append("" if url is None else str(url).strip())
            return (
                "" if channel is None else str(channel).strip(),
                rows,
                urls,
            )
        finally:
            workbook.close()

    def merge_group(self, group: dict, output_path) -> None:

        inputs = [Path(item) for item in group["input_files"]]
        if not inputs:
            raise AdapterError("KYD 合并分组没有输入文件")

        output_path = Path(output_path)
        if output_path.suffix.lower() != ".xls":
            output_path = output_path.with_suffix(".xls")
        output_path.parent.mkdir(parents=True, exist_ok=True)

        copy_template(KYD_TEMPLATE_PATH, output_path)

        channel = ""
        lines: list[KydLine] = []

        for source in inputs:
            src_channel, rows, urls = self._source_rows(source)
            if src_channel and not channel:
                channel = src_channel
            if not rows:
                raise AdapterError(f"{source.name}：没有找到明细行")
            if len(rows) != len(urls):
                raise AdapterError(f"{source.name}：明细与图片网址数量不一致")
            for values, url in zip(rows, urls):
                image, kind = download_product_image(url, source.name)
                lines.append(
                    KydLine(
                        cells=values,
                        image_bytes=image,
                        image_kind=kind,
                    )
                )

        if not channel:
            raise AdapterError("源发票没有渠道（B4）")

        fill_kyd_template(
            output_path,
            channel=channel,
            warehouse=str(group["warehouse_code"]).strip(),
            carton_count=len(group["carton_numbers"]),
            lines=lines,
        )

    def validate_output(self, output_path, group: dict) -> list[str]:

        errors = []
        path = Path(output_path)
        warehouse = str(group["warehouse_code"]).strip()

        if not path.exists():
            return [f"输出文件不存在：{path}"]
        if path.suffix.lower() != ".xls":
            return [f"{path.name}：KYD 输出必须是 .xls"]
        if path.name.startswith("~$"):
            return [f"输出是临时锁文件：{path.name}"]

        try:
            import olefile
            import xlrd
        except ImportError as exc:
            return [f"{path.name}：缺少校验依赖：{exc}"]

        ole = olefile.OleFileIO(str(path))
        try:
            streams = ["/".join(item) for item in ole.listdir()]
            if any("CompObj" in name for name in streams):
                errors.append(f"{path.name}：被 Excel 另存，快越达无法识别图片")
            book = ole.openstream(["Workbook"]).read()
        finally:
            ole.close()

        try:
            wb = xlrd.open_workbook(str(path), on_demand=True)
        except Exception as exc:
            return [f"{path.name}：无法打开：{exc}"]

        try:
            names = wb.sheet_names()
            for required in (self.OUTPUT_SHEET, "服务", "地址库"):
                if required not in names:
                    errors.append(f"{path.name}：缺少工作表【{required}】")
            if errors:
                return errors

            ws = wb.sheet_by_name(self.OUTPUT_SHEET)
            actual_wh = str(ws.cell_value(4, 1)).strip()
            if actual_wh.upper() != warehouse.upper():
                errors.append(
                    f"{path.name}：仓库单元格不是 {warehouse}"
                )

            try:
                carton_value = int(float(ws.cell_value(25, 1)))
            except (TypeError, ValueError):
                carton_value = None
            if carton_value != len(group["carton_numbers"]):
                errors.append(
                    f"{path.name}：B26 应为 {len(group['carton_numbers'])}"
                )

            data_rows = 0
            for row in range(29, ws.nrows):
                carton = ws.cell_value(row, 0)
                product = ws.cell_value(row, 2)
                if carton not in ("", None) or product not in ("", None):
                    data_rows += 1
            if data_rows < group["input_count"]:
                errors.append(
                    f"{path.name}：明细行 {data_rows}，"
                    f"少于箱数 {group['input_count']}"
                )

            try:
                address = lookup_kyd_address(warehouse)
                zip_value = ws.cell_value(11, 1)
                if address.zip_code is not None and _is_number(zip_value):
                    if abs(float(zip_value) - float(address.zip_code)) >= 0.1:
                        errors.append(
                            f"{path.name}：邮编 {zip_value} "
                            f"与地址库 {address.zip_code} 不匹配"
                        )
            except AdapterError as exc:
                errors.append(str(exc))
        finally:
            wb.release_resources()

        pic_count = 0
        for _pos, rec, ln, payload in iter_biff(book):
            if rec == 0x005D and ln >= 8:
                _ft, _cb, ot, _oid = struct.unpack_from(
                    "<HHHH",
                    payload,
                    0,
                )
                if ot == 8:
                    pic_count += 1
        if pic_count < group["input_count"]:
            errors.append(
                f"{path.name}：浮动图片 {pic_count} 张，"
                f"少于箱数 {group['input_count']}"
            )

        return errors


def _is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _blank(value) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def prefer_jpeg_url(url: str) -> str:
    """
    Amazon 部分地址带 FMwebp，后缀虽是 .jpg，实际返回 WebP。
    去掉 FMwebp 后通常能拿到真正的 JPEG。
    """

    text = (url or "").strip()
    if not text:
        return text
    return (
        text.replace("_FMwebp_", "_")
        .replace("FMwebp_", "")
        .replace("_FMwebp", "")
        .replace("FMwebp", "")
    )


def detect_image_kind(data: bytes) -> str:
    if data.startswith(b"\xff\xd8"):
        return "jpeg"
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if len(data) >= 12 and data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "webp"
    if data.startswith(b"GIF8"):
        return "gif"
    if data.startswith(b"BM"):
        return "bmp"
    head = data.lstrip()[:32].lower()
    if head.startswith(b"<") or head.startswith(b"<!doctype"):
        raise AdapterError("产品图片网址返回了网页，不是图片")
    from io import BytesIO

    from PIL import Image

    try:
        Image.open(BytesIO(data))
    except Exception as exc:
        raise AdapterError(
            "产品图片不是可识别的图片格式"
        ) from exc
    return "other"


def _http_get_image(url: str) -> bytes:
    request = Request(
        url,
        headers={"User-Agent": "Mozilla/5.0 InvoiceMergeSystem"},
    )
    with urlopen(request, timeout=30) as response:
        return response.read()


def download_product_image(url: str, source_name: str) -> tuple[bytes, str]:

    if not url or not url.lower().startswith(("http://", "https://")):
        raise AdapterError(f"{source_name}：产品图片网址为空")

    candidates = []
    jpeg_url = prefer_jpeg_url(url)
    if jpeg_url != url:
        candidates.append(jpeg_url)
    candidates.append(url)

    data = b""
    last_error: Exception | None = None
    for candidate in candidates:
        try:
            data = _http_get_image(candidate)
            if data:
                break
        except URLError as exc:
            last_error = exc
            data = b""

    if not data:
        if last_error is not None:
            raise AdapterError(
                f"{source_name}：无法下载产品图：{url}：{last_error}"
            ) from last_error
        raise AdapterError(f"{source_name}：产品图下载结果为空")

    try:
        return data, detect_image_kind(data)
    except AdapterError as exc:
        raise AdapterError(f"{source_name}：{exc}") from exc

