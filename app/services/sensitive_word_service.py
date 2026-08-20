from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from app.invoice_config import SYSTEM_ROOT
from app.json_io import read_json, write_json
from app.services.invoice_adapters import get_adapter
from app.services.invoice_scan_service import collect_invoice_files


SENSITIVE_WORDS_PATH = (
    SYSTEM_ROOT / "公共配置" / "sensitive_words.json"
)

FIELD_ALIASES = {
    "ProductEN": ["产品英文品名", "产品英文品名*"],
    "ProductCN": ["产品中文品名", "产品中文品名*"],
    "MaterialEN": ["英文产品材质", "产品材质", "产品材质*"],
    "MaterialCN": ["中文产品材质", "产品材质", "产品材质*"],
}

FIELD_NAMES = {
    "ProductEN": "产品英文品名",
    "ProductCN": "产品中文品名",
    "MaterialEN": "英文产品材质",
    "MaterialCN": "中文产品材质",
}

WHITESPACE_RE = re.compile(r"\s+")
INVISIBLE_RE = re.compile(r"[\u200b\u200c\u200d\ufeff\u00a0\u202a-\u202e]")
FORMULA_STRING_RE = re.compile(r'"([^"]*)"')


@dataclass
class KeywordRule:
    keyword: str
    category: str
    enabled: bool
    fields: list[str]
    whitelist_phrases: list[str]
    remark: str = ""


@dataclass
class SensitiveHit:
    file_path: str
    sheet_name: str
    cell_address: str
    field_code: str
    field_name: str
    keyword: str
    category: str
    original_text: str
    normalized_text: str


@dataclass
class SensitiveScanResult:
    file_count: int
    field_count: int
    hit_count: int
    hit_file_count: int
    hits: list[SensitiveHit] = field(default_factory=list)
    aggregates: list[dict] = field(default_factory=list)
    warning_messages: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def cell_to_text(value) -> str:

    if value is None:
        return ""

    try:
        from openpyxl.cell.rich_text import CellRichText, TextBlock

        if isinstance(value, CellRichText):
            parts = []
            for part in value:
                if isinstance(part, str):
                    parts.append(part)
                elif isinstance(part, TextBlock):
                    parts.append(part.text or "")
                else:
                    parts.append(str(part))
            return "".join(parts)
    except Exception:
        pass

    if isinstance(value, str) and value.startswith("="):
        literals = FORMULA_STRING_RE.findall(value)
        return " ".join(literals)

    return str(value)


def normalize_text(value) -> str:

    text = cell_to_text(value)
    if not text:
        return ""

    text = INVISIBLE_RE.sub("", text)
    text = text.replace("\r\n", " ").replace("\n", " ").replace("\r", " ")
    text = WHITESPACE_RE.sub(" ", text).strip()
    return text


def apply_whitelist(text: str, phrases: list[str], ignore_case: bool) -> str:

    result = text

    for phrase in phrases:
        if not phrase:
            continue
        if ignore_case:
            result = re.sub(
                re.escape(phrase),
                "",
                result,
                flags=re.IGNORECASE,
            )
        else:
            result = result.replace(phrase, "")

    return normalize_text(result) if phrases else result


def keyword_hits_in_text(
    text: str,
    keyword: str,
    whitelist_phrases: list[str],
) -> bool:

    if not text or not keyword:
        return False

    text = INVISIBLE_RE.sub("", text)
    keyword = INVISIBLE_RE.sub("", keyword)
    has_latin = bool(re.search(r"[A-Za-z]", keyword))
    prepared = apply_whitelist(text, whitelist_phrases, ignore_case=has_latin)

    if has_latin:
        return keyword.lower() in prepared.lower()

    compact_text = WHITESPACE_RE.sub("", prepared)
    compact_keyword = WHITESPACE_RE.sub("", keyword)
    return compact_keyword in compact_text


@dataclass
class ScanProfile:
    carrier_code: str
    carrier_name: str
    data_start_row: int
    header_row: int
    sheet_name: str


EMPTY_CARTON_STOP = 8


def default_scan_profiles() -> list[dict]:

    return [
        {
            "carrier_code": "MC",
            "carrier_name": "迈创",
            "data_start_row": 18,
            "header_row": 17,
            "sheet_name": "迈创发票",
        },
        {
            "carrier_code": "KYD",
            "carrier_name": "快越达",
            "data_start_row": 30,
            "header_row": 29,
            "sheet_name": "快越达发票",
        },
    ]


def default_config() -> dict:

    return {
        "version": "1.1",
        "keywords": [],
        "scan_profiles": default_scan_profiles(),
    }


def _merge_profiles(raw_profiles) -> list[dict]:

    merged = {
        item["carrier_code"]: dict(item)
        for item in default_scan_profiles()
    }

    if isinstance(raw_profiles, list):
        for item in raw_profiles:
            if not isinstance(item, dict):
                continue
            code = str(item.get("carrier_code", "")).strip().upper()
            if not code:
                continue
            current = merged.get(code, {"carrier_code": code})
            current["carrier_code"] = code
            if item.get("carrier_name"):
                current["carrier_name"] = str(item["carrier_name"]).strip()
            if item.get("sheet_name"):
                current["sheet_name"] = str(item["sheet_name"]).strip()
            try:
                start = int(item.get("data_start_row") or 0)
                if start >= 1:
                    current["data_start_row"] = start
            except (TypeError, ValueError):
                pass
            try:
                header = int(item.get("header_row") or 0)
                if header >= 1:
                    current["header_row"] = header
            except (TypeError, ValueError):
                pass
            merged[code] = current

    return [merged[code] for code in sorted(merged)]


def load_full_config(path: Path | None = None) -> dict:

    config_path = path or SENSITIVE_WORDS_PATH

    if not config_path.exists():
        data = default_config()
        write_json(config_path, data)
        return data

    data = read_json(config_path)
    if not isinstance(data, dict):
        data = default_config()
        write_json(config_path, data)
        return data

    original_version = str(data.get("version") or "")
    original_profiles = data.get("scan_profiles")
    data["version"] = "1.1"
    data["keywords"] = list(data.get("keywords") or [])
    data["scan_profiles"] = _merge_profiles(original_profiles)
    if not original_profiles or original_version != "1.1":
        write_json(config_path, data)
    return data


def save_full_config(data: dict, path: Path | None = None) -> Path:

    config_path = path or SENSITIVE_WORDS_PATH
    payload = {
        "version": str(data.get("version") or "1.1"),
        "keywords": list(data.get("keywords") or []),
        "scan_profiles": _merge_profiles(data.get("scan_profiles")),
    }
    write_json(config_path, payload)
    return config_path


def load_scan_profiles(path: Path | None = None) -> dict[str, ScanProfile]:

    data = load_full_config(path)
    result = {}

    for item in data.get("scan_profiles", []):
        code = str(item.get("carrier_code", "")).strip().upper()
        if not code:
            continue
        result[code] = ScanProfile(
            carrier_code=code,
            carrier_name=str(item.get("carrier_name", "")).strip(),
            data_start_row=int(item.get("data_start_row") or 1),
            header_row=int(item.get("header_row") or 1),
            sheet_name=str(item.get("sheet_name", "")).strip(),
        )

    return result


def load_keyword_rules(path: Path | None = None) -> list[KeywordRule]:

    data = load_full_config(path)
    rules = []

    for item in data.get("keywords", []):
        rules.append(
            KeywordRule(
                keyword=str(item.get("keyword", "")).strip(),
                category=str(item.get("category", "")).strip(),
                enabled=bool(item.get("enabled", True)),
                fields=list(item.get("fields") or list(FIELD_NAMES)),
                whitelist_phrases=[
                    str(phrase)
                    for phrase in item.get("whitelist_phrases", [])
                    if str(phrase).strip()
                ],
                remark=str(item.get("remark", "")),
            )
        )

    return [rule for rule in rules if rule.enabled and rule.keyword]


def detect_carrier(path: Path, workbook) -> str:

    meta = _meta_from_sheet(workbook, "_SystemMeta")
    if not meta:
        meta = _meta_from_sheet(workbook, "_MergeMeta")

    code = str(meta.get("CarrierCode", "")).strip().upper()
    if code:
        return code

    names = set(workbook.sheetnames)
    if "迈创发票" in names:
        return "MC"
    if "快越达发票" in names:
        return "KYD"

    haystack = str(path)
    name_upper = path.name.upper()
    if "迈创" in haystack or name_upper.startswith("MC_") or "_MC_" in f"_{name_upper}_":
        return "MC"
    if "快越达" in haystack or name_upper.startswith("KYD_") or "_KYD_" in f"_{name_upper}_":
        return "KYD"

    return ""


META_SHEET_NAMES = {"_SystemMeta", "_MergeMeta"}
FALLBACK_SHEET_NAMES = ("合并发票",)


def _candidate_sheet_names(workbook, preferred: str = "") -> list[str]:

    names = []
    if preferred and preferred in workbook.sheetnames:
        names.append(preferred)
    for extra in FALLBACK_SHEET_NAMES:
        if extra in workbook.sheetnames and extra not in names:
            names.append(extra)
    for name in workbook.sheetnames:
        if name not in META_SHEET_NAMES and name not in names:
            names.append(name)
    return names


def _row_headers(worksheet, row_number: int) -> list[str]:

    first = next(
        worksheet.iter_rows(
            min_row=row_number,
            max_row=row_number,
            values_only=True,
        ),
        None,
    )
    if not first:
        return []
    return [str(value).strip() if value is not None else "" for value in first]


def _locate_header(
    worksheet,
    preferred_row: int | None = None,
) -> tuple[int, dict[str, int]] | None:

    last = min(worksheet.max_row or 1, 40)
    candidates = []
    if preferred_row and preferred_row >= 1:
        candidates.append(preferred_row)
    candidates.extend(row for row in range(1, last + 1) if row not in candidates)

    for row in candidates:
        mapping = _header_map(_row_headers(worksheet, row))
        if "ProductCN" in mapping or "ProductEN" in mapping:
            return row, mapping
    return None


def _cells_from_mapped_sheet(
    worksheet,
    sheet_name: str,
    header_row: int,
    start_row: int,
    mapping: dict[str, int],
) -> list[tuple[str, str, str, str, str]]:

    from openpyxl.utils import get_column_letter

    results = []
    data_start = max(start_row, header_row + 1)
    field_columns = [index + 1 for index in mapping.values()]

    for row_number in _iter_data_row_numbers(worksheet, data_start, field_columns):
        for field_code, index in mapping.items():
            cell = worksheet.cell(row_number, index + 1)
            results.append(
                (
                    sheet_name,
                    f"{get_column_letter(index + 1)}{row_number}",
                    field_code,
                    FIELD_NAMES[field_code],
                    cell.value,
                )
            )
    return results


def _iter_data_row_numbers(worksheet, start_row: int, columns: list[int] | None = None) -> list[int]:

    rows = []
    empty_run = 0
    last = min(worksheet.max_row or start_row, start_row + 400)
    watch_cols = [1]
    if columns:
        watch_cols.extend(col for col in columns if col >= 1 and col not in watch_cols)

    for row in range(start_row, last + 1):
        occupied = False
        for col in watch_cols:
            if normalize_text(worksheet.cell(row, col).value):
                occupied = True
                break
        if occupied:
            rows.append(row)
            empty_run = 0
        else:
            empty_run += 1
            if empty_run >= EMPTY_CARTON_STOP:
                break

    return rows


def _norm_header(text: str) -> str:

    return normalize_text(text).replace("*", "").replace(" ", "")


def _header_map(headers: list[str]) -> dict[str, int]:

    mapping = {}
    normalized = [_norm_header(item) for item in headers]

    for field_code, aliases in FIELD_ALIASES.items():
        alias_norms = [_norm_header(alias) for alias in aliases if alias]
        for index, header in enumerate(normalized):
            if header and header in alias_norms:
                mapping[field_code] = index
                break
        if field_code in mapping:
            continue
        for index, header in enumerate(normalized):
            if header and any(header.startswith(alias) for alias in alias_norms):
                mapping[field_code] = index
                break

    return mapping


def _meta_from_sheet(workbook, sheet_name: str) -> dict:

    if sheet_name not in workbook.sheetnames:
        return {}

    values = {}
    ws = workbook[sheet_name]

    for row in ws.iter_rows(min_col=1, max_col=2, values_only=True):
        if row[0] is None:
            continue
        key = str(row[0]).strip()
        values[key] = "" if row[1] is None else str(row[1]).strip()

    return values


def _cells_from_adapter(
    path: Path,
    workbook,
    start_row_override: int | None = None,
) -> list[tuple[str, str, str, str, str]]:

    meta = _meta_from_sheet(workbook, "_SystemMeta")
    if not meta:
        meta = _meta_from_sheet(workbook, "_MergeMeta")

    carrier = str(meta.get("CarrierCode", "")).strip().upper()
    if not carrier:
        carrier = detect_carrier(path, workbook)

    version = meta.get("TemplateVersion") or "1.0"
    if version == "1":
        version = "1.0"

    if not carrier:
        return []

    try:
        adapter = get_adapter(carrier, version)
    except Exception:
        return []

    cfg = adapter.template_config()
    start_row = start_row_override or cfg.data_start_row
    if start_row < 1:
        return []

    for sheet_name in _candidate_sheet_names(workbook, cfg.sheet_name):
        ws = workbook[sheet_name]
        results = []
        field_columns = [
            spec.column
            for spec in adapter.product_field_specs()
            if spec.column >= 1
        ]
        for spec in adapter.product_field_specs():
            if spec.column < 1:
                continue
            for row in _iter_data_row_numbers(ws, start_row, field_columns):
                cell = ws.cell(row, spec.column)
                results.append(
                    (
                        sheet_name,
                        cell.coordinate,
                        spec.field_code,
                        spec.field_name,
                        cell.value,
                    )
                )
        if results:
            return results

    return []


def _cells_from_profile(
    workbook,
    profile: ScanProfile,
) -> list[tuple[str, str, str, str, str]]:

    header_row = max(profile.header_row, 1)
    start_row = max(profile.data_start_row, header_row + 1)

    for sheet_name in _candidate_sheet_names(workbook, profile.sheet_name):
        ws = workbook[sheet_name]
        located = _locate_header(ws, header_row)
        if not located:
            continue
        found_header, mapping = located
        results = _cells_from_mapped_sheet(
            ws,
            sheet_name,
            found_header,
            start_row if found_header == header_row else found_header + 1,
            mapping,
        )
        if results:
            return results

    return []


def _cells_from_headers(workbook) -> list[tuple[str, str, str, str, str]]:

    results = []

    for sheet_name in _candidate_sheet_names(workbook):
        ws = workbook[sheet_name]
        located = _locate_header(ws)
        if not located:
            continue
        header_row, mapping = located
        results.extend(
            _cells_from_mapped_sheet(
                ws,
                sheet_name,
                header_row,
                header_row + 1,
                mapping,
            )
        )

    return results


def _scan_kyd_xls(path: Path, rules: list[KeywordRule]) -> list[SensitiveHit]:

    import xlrd
    from openpyxl.utils import get_column_letter

    workbook = xlrd.open_workbook(str(path))
    names = workbook.sheet_names()
    sheet_name = "模板" if "模板" in names else names[0]
    worksheet = workbook.sheet_by_name(sheet_name)
    specs = [
        ("ProductEN", FIELD_NAMES["ProductEN"], 2),
        ("ProductCN", FIELD_NAMES["ProductCN"], 3),
        ("MaterialEN", FIELD_NAMES["MaterialEN"], 12),
        ("MaterialCN", FIELD_NAMES["MaterialCN"], 12),
    ]

    cells = []
    for row in range(29, worksheet.nrows):
        carton = worksheet.cell_value(row, 0)
        product = worksheet.cell_value(row, 2)
        if carton in ("", None) and product in ("", None):
            continue
        for field_code, field_name, col in specs:
            cells.append(
                (
                    sheet_name,
                    f"{get_column_letter(col + 1)}{row + 1}",
                    field_code,
                    field_name,
                    worksheet.cell_value(row, col),
                )
            )

    hits: list[SensitiveHit] = []
    for sheet_name, address, field_code, field_name, raw in cells:
        original = cell_to_text(raw)
        normalized = normalize_text(raw)
        for rule in rules:
            if field_code not in rule.fields:
                continue
            if keyword_hits_in_text(
                normalized,
                rule.keyword,
                rule.whitelist_phrases,
            ):
                hits.append(
                    SensitiveHit(
                        file_path=str(path),
                        sheet_name=sheet_name,
                        cell_address=address,
                        field_code=field_code,
                        field_name=field_name,
                        keyword=rule.keyword,
                        category=rule.category,
                        original_text=original,
                        normalized_text=normalized,
                    )
                )
    return hits


def scan_workbook(
    path: Path,
    rules: list[KeywordRule],
    profiles: dict[str, ScanProfile] | None = None,
) -> list[SensitiveHit]:

    if path.suffix.lower() == ".xls":
        return _scan_kyd_xls(path, rules)

    workbook = load_workbook(path, data_only=False, read_only=False)
    hits: list[SensitiveHit] = []

    try:
        profiles = profiles if profiles is not None else load_scan_profiles()
        carrier = detect_carrier(path, workbook)
        profile = profiles.get(carrier)
        start_override = profile.data_start_row if profile else None

        cells = []
        if profile:
            cells = _cells_from_profile(workbook, profile)
        if not cells:
            cells = _cells_from_adapter(path, workbook, start_override)
        if not cells:
            cells = _cells_from_headers(workbook)

        for sheet_name, address, field_code, field_name, raw in cells:
            original = cell_to_text(raw)
            normalized = normalize_text(raw)

            for rule in rules:
                if field_code not in rule.fields:
                    continue

                if keyword_hits_in_text(
                    normalized,
                    rule.keyword,
                    rule.whitelist_phrases,
                ):
                    hits.append(
                        SensitiveHit(
                            file_path=str(path),
                            sheet_name=sheet_name,
                            cell_address=address,
                            field_code=field_code,
                            field_name=field_name,
                            keyword=rule.keyword,
                            category=rule.category,
                            original_text=original,
                            normalized_text=normalized,
                        )
                    )
    finally:
        workbook.close()

    return hits


def aggregate_hits(hits: list[SensitiveHit]) -> list[dict]:

    grouped: dict[tuple[str, str], dict] = {}

    for hit in hits:
        key = (hit.file_path, hit.keyword)
        item = grouped.get(key)
        if item is None:
            item = {
                "file_path": hit.file_path,
                "file_name": Path(hit.file_path).name,
                "keyword": hit.keyword,
                "category": hit.category,
                "count": 0,
                "details": [],
                "cell_labels": [],
            }
            grouped[key] = item

        item["count"] += 1
        item["details"].append(hit)
        item["cell_labels"].append(f"{hit.sheet_name}!{hit.cell_address}")

    return list(grouped.values())


def scan_paths(
    paths: list[Path],
    *,
    recursive: bool = True,
    config_path: Path | None = None,
) -> SensitiveScanResult:

    rules = load_keyword_rules(config_path)
    profiles = load_scan_profiles(config_path)
    files: list[Path] = []

    for path in paths:
        if path.is_file() and path.suffix.lower() in {".xlsx", ".xls"} and not path.name.startswith("~$"):
            files.append(path)
        elif path.is_dir():
            if recursive:
                files.extend(collect_invoice_files([path]))
            else:
                files.extend(
                    sorted(
                        item
                        for pattern in ("*.xlsx", "*.xls")
                        for item in path.glob(pattern)
                        if not item.name.startswith("~$")
                    )
                )

    unique = []
    seen = set()
    for file in files:
        resolved = file.resolve()
        if resolved in seen:
            continue
        seen.add(resolved)
        unique.append(file)

    hits: list[SensitiveHit] = []
    errors: list[str] = []
    field_count = 0
    warnings: list[str] = []

    if not rules:
        warnings.append(
            "敏感词词库为空，请先维护词库后再检查"
        )

    for file in unique:
        try:
            file_hits = scan_workbook(file, rules, profiles)
            hits.extend(file_hits)
            field_count += 4
        except Exception as exc:
            errors.append(f"{file.name}：{exc}")

    aggregates = aggregate_hits(hits)
    hit_files = {hit.file_path for hit in hits}

    warnings.extend(
        f"{Path(item['file_path']).name} / {item['keyword']} x{item['count']}"
        for item in aggregates
    )

    return SensitiveScanResult(
        file_count=len(unique),
        field_count=field_count,
        hit_count=len(hits),
        hit_file_count=len(hit_files),
        hits=hits,
        aggregates=aggregates,
        warning_messages=warnings,
        errors=errors,
    )


def scan_output_files(files: list[Path]) -> SensitiveScanResult:

    return scan_paths(files, recursive=False)


def export_results(result: SensitiveScanResult, output_path: Path) -> Path:

    from openpyxl import Workbook

    workbook = Workbook()
    summary = workbook.active
    summary.title = "汇总"
    summary.append(["文件", "敏感词", "分类", "单元格", "命中次数"])

    for item in result.aggregates:
        summary.append(
            [
                item["file_path"],
                item["keyword"],
                item["category"],
                "、".join(item.get("cell_labels") or []),
                item["count"],
            ]
        )

    detail = workbook.create_sheet("明细")
    detail.append(
        [
            "文件",
            "工作表",
            "单元格",
            "字段代码",
            "字段名",
            "敏感词",
            "分类",
            "原始文本",
            "规范化文本",
        ]
    )

    for hit in result.hits:
        detail.append(
            [
                hit.file_path,
                hit.sheet_name,
                hit.cell_address,
                hit.field_code,
                hit.field_name,
                hit.keyword,
                hit.category,
                hit.original_text,
                hit.normalized_text,
            ]
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return output_path
