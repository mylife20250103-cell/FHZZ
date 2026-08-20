from __future__ import annotations

import shutil
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from app.invoice_config import LOCAL_TEMP_ROOT
from app.services.invoice_scan_service import collect_invoice_files, read_system_meta


@dataclass
class ModifyPreviewRow:
    path: str
    file_name: str
    carrier_code: str
    sheet_name: str
    cell: str
    current_value: str
    new_value: str
    selected: bool = True
    status: str = "待处理"
    error: str = ""


@dataclass
class ModifyResult:
    passed: bool
    success_count: int
    fail_count: int
    rows: list[ModifyPreviewRow]
    errors: list[str] = field(default_factory=list)


def normalize_cell_address(address: str) -> str:
    """
    把用户输入收成 Excel 地址，例如：
    B2、b2、B2单元格、单元格B2 → B2
    """

    text = (address or "").strip().upper()
    text = (
        text.replace("单元格", "")
        .replace("CELL", "")
        .replace("$", "")
        .replace("：", "")
        .replace(":", "")
        .strip()
    )

    letters = ""
    digits = ""

    for char in text:
        if "A" <= char <= "Z":
            if digits:
                break
            letters += char
        elif char.isdigit():
            if not letters:
                continue
            digits += char

    if not letters or not digits:
        return ""

    return f"{letters}{digits}"


def parse_cell(address: str) -> tuple[str, int]:

    text = normalize_cell_address(address)

    if not text:
        raise ValueError(f"无效单元格：{address}")

    letters = "".join(
        char for char in text if "A" <= char <= "Z"
    )
    digits = "".join(
        char for char in text if char.isdigit()
    )

    if not letters or not digits:
        raise ValueError(f"无效单元格：{address}")

    return letters, int(digits)


def preview_files(
    paths: list[Path],
    new_value: str,
    manual_cell: str = "",
    recursive: bool = True,
) -> list[ModifyPreviewRow]:

    files: list[Path] = []

    for path in paths:
        if path.is_file() and path.suffix.lower() == ".xlsx":
            files.append(path)
        elif path.is_dir():
            if recursive:
                files.extend(collect_invoice_files([path]))
            else:
                files.extend(
                    item
                    for item in path.glob("*.xlsx")
                    if not item.name.startswith("~$")
                )

    rows = []

    for path in files:
        if path.name.startswith("~$"):
            continue

        carrier = ""
        cell = normalize_cell_address(manual_cell)
        sheet_name = ""
        current = ""
        error = ""

        meta, meta_error = read_system_meta(path)

        if meta:
            carrier = meta.get("CarrierCode", "")
            cell = (
                normalize_cell_address(meta.get("ChannelCell", ""))
                or cell
            )
            sheet_name = next(
                (
                    name
                    for name in load_sheet_names(path)
                    if name not in {"_SystemMeta", "_MergeMeta"}
                ),
                "",
            )
        elif meta_error and "缺少 _SystemMeta" not in meta_error:
            error = meta_error

        if not sheet_name:
            names = load_sheet_names(path)
            sheet_name = next(
                (
                    name
                    for name in names
                    if name not in {"_SystemMeta", "_MergeMeta"}
                ),
                names[0] if names else "",
            )

        if cell and sheet_name and not error:
            try:
                current = read_cell_value(path, sheet_name, cell)
            except Exception as exc:
                error = str(exc)

        rows.append(
            ModifyPreviewRow(
                path=str(path),
                file_name=path.name,
                carrier_code=carrier,
                sheet_name=sheet_name,
                cell=cell,
                current_value=current,
                new_value=new_value,
                selected=True,
                status="待处理" if not error else "预检失败",
                error=error,
            )
        )

    return rows


def load_sheet_names(path: Path) -> list[str]:

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def read_cell_value(path: Path, sheet_name: str, cell: str) -> str:

    workbook = load_workbook(path, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"找不到工作表：{sheet_name}")
        value = workbook[sheet_name][cell].value
        return "" if value is None else str(value)
    finally:
        workbook.close()


def precheck_rows(rows: list[ModifyPreviewRow]) -> list[str]:

    errors = []

    for row in rows:
        if not row.selected:
            continue

        path = Path(row.path)

        if not path.exists():
            errors.append(f"{row.file_name}：文件不存在")
            continue

        lock = path.with_name(f"~${path.name[2:] if path.name.startswith('~$') else path.name}")
        # Excel lock file
        lock_file = path.parent / f"~${path.name}"
        if lock_file.exists():
            errors.append(f"{row.file_name}：文件可能已被 Excel 打开")

        if not row.cell:
            errors.append(f"{row.file_name}：未指定目标单元格")
            continue

        try:
            parse_cell(row.cell)
        except Exception as exc:
            errors.append(f"{row.file_name}：{exc}")
            continue

        if not row.sheet_name:
            errors.append(f"{row.file_name}：未找到目标工作表")

    return errors


def apply_modifications(rows: list[ModifyPreviewRow]) -> ModifyResult:

    selected = [row for row in rows if row.selected]
    errors = precheck_rows(selected)

    if errors:
        return ModifyResult(
            passed=False,
            success_count=0,
            fail_count=len(selected),
            rows=rows,
            errors=errors,
        )

    backup_root = LOCAL_TEMP_ROOT / "batch_modify_backup"
    if backup_root.exists():
        shutil.rmtree(backup_root, ignore_errors=True)
    backup_root.mkdir(parents=True, exist_ok=True)

    changed: list[tuple[ModifyPreviewRow, Path]] = []

    try:
        from app.excel_com import (
            close_workbook,
            excel_application,
            open_workbook,
        )

        with excel_application() as excel:
            for row in selected:
                path = Path(row.path)
                backup = backup_root / path.name
                shutil.copy2(path, backup)

                workbook = open_workbook(excel, path, read_only=False)
                try:
                    ws = workbook.Worksheets(row.sheet_name)
                    ws.Range(row.cell).Value = row.new_value
                    workbook.Save()
                finally:
                    close_workbook(workbook, save=True)

                actual = read_cell_value(path, row.sheet_name, row.cell)
                if actual != str(row.new_value):
                    raise ValueError(
                        f"{row.file_name} 回读校验失败："
                        f"期望 {row.new_value}，实际 {actual}"
                    )

                row.status = "成功"
                row.current_value = actual
                changed.append((row, backup))

        shutil.rmtree(backup_root, ignore_errors=True)

        return ModifyResult(
            passed=True,
            success_count=len(selected),
            fail_count=0,
            rows=rows,
            errors=[],
        )

    except Exception as exc:
        for row, backup in reversed(changed):
            try:
                shutil.copy2(backup, row.path)
                row.status = "已回滚"
            except Exception:
                row.status = "回滚失败"

        for row in selected:
            if row.status == "待处理":
                row.status = "失败"
                row.error = str(exc)

        return ModifyResult(
            passed=False,
            success_count=0,
            fail_count=len(selected),
            rows=rows,
            errors=[str(exc)],
        )
