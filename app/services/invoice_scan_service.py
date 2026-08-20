from __future__ import annotations

import configparser
import hashlib
import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from app.invoice_config import (
    ALLOWED_EXTENSIONS,
    CURRENT_CONFIG_INI,
    EXCLUDE_DIR_NAMES,
    INVOICE_SOURCES_INI,
    is_excel_junk_file,
)


SOURCE_ID_PATTERN = re.compile(r"^[0-9A-F]{8}$")
CARTON_SUFFIX_PATTERN = re.compile(r"^\d{6}$")


META_FIELDS = [
    "MetaSchemaVersion",
    "CarrierCode",
    "CarrierName",
    "TemplateVersion",
    "InvoiceType",
    "ChannelCell",
    "DateID",
    "GeneratedAt",
    "StoreCode",
    "PlanID",
    "SourceID",
    "WarehouseCode",
    "FBABatch",
    "CartonNumber",
]


@dataclass
class InvoiceRecord:
    path: str
    sha256: str

    carrier_code: str
    carrier_name: str
    template_version: str

    date_id: str
    generated_at: str

    store_code: str
    plan_id: str
    source_id: str

    warehouse_code: str
    fba_batch: str
    carton_number: str


@dataclass(frozen=True)
class RegisteredSource:
    source_key: str
    path: Path
    owner: str
    store_code: str
    carrier_code: str


@dataclass
class ScanResult:
    passed: bool

    invoice_count: int
    carton_count: int
    source_count: int
    warehouse_count: int
    carrier_count: int

    records: list[InvoiceRecord]
    errors: list[str]
    warnings: list[str]

    batch_id: str | None = None
    snapshot_path: str | None = None


# =========================================================
# SHA256
# =========================================================

def sha256_file(path: Path) -> str:

    digest = hashlib.sha256()

    with path.open("rb") as file:

        for chunk in iter(
            lambda: file.read(
                1024 * 1024
            ),
            b"",
        ):
            digest.update(chunk)

    return digest.hexdigest()


# =========================================================
# INI
# =========================================================

def load_ini(path: Path):

    parser = configparser.ConfigParser(
        interpolation=None
    )

    parser.optionxform = str

    encodings = [
        "utf-8-sig",
        "gbk",
        "utf-8",
    ]

    last_error = None

    for encoding in encodings:

        try:

            with path.open(
                "r",
                encoding=encoding,
            ) as f:

                parser.read_file(f)

            return parser

        except UnicodeDecodeError as exc:

            parser.clear()
            last_error = exc

    if last_error:
        raise last_error

    return parser


def load_known_carriers() -> set[str]:

    if not CURRENT_CONFIG_INI.exists():

        raise FileNotFoundError(
            "中央配置不存在：\n"
            f"{CURRENT_CONFIG_INI}"
        )

    parser = load_ini(
        CURRENT_CONFIG_INI
    )

    result = set()

    for section in parser.sections():

        match = re.fullmatch(
            r"Carrier\.([A-Za-z0-9_-]+)",
            section,
        )

        if match:
            result.add(
                match.group(1).upper()
            )

    if not result:

        raise ValueError(
            "current_config.ini "
            "没有任何 [Carrier.xxx] 配置"
        )

    return result


def is_enabled_flag(value: str) -> bool:

    return value.strip().upper() in {
        "1",
        "TRUE",
        "YES",
        "Y",
        "ON",
    }


def validate_source_id(source_id: str) -> str | None:

    if not SOURCE_ID_PATTERN.fullmatch(source_id or ""):
        return (
            "SourceID 必须为8位大写十六进制："
            f"{source_id}"
        )

    return None


def validate_carton_number(
    fba_batch: str,
    carton: str,
) -> str | None:

    prefix = f"{fba_batch}U"

    if not carton.startswith(prefix):
        return (
            "CartonNumber 与 FBABatch 不匹配："
            f"{carton}"
        )

    suffix = carton[len(prefix):]

    if not CARTON_SUFFIX_PATTERN.fullmatch(suffix):
        return "CartonNumber 后缀必须为6位数字"

    if int(suffix) < 1:
        return "CartonNumber 不允许 U000000"

    return None


def machine_key(
    source_id: str,
    carton_number: str,
) -> tuple[str, str]:

    return (source_id, carton_number)


# =========================================================
# 扫描源
# =========================================================

def load_source_entries() -> list[RegisteredSource]:
    """
    读取 invoice_sources.ini 中已启用的 [Source.xxx]。

    每个物流商一条扫描源。发票中心一次只选其中一个再扫描。
    Enabled 支持：1 / TRUE / YES / Y / ON（大小写不敏感）。
    如果文件还没建立，返回空列表。
    """

    if not INVOICE_SOURCES_INI.exists():
        return []

    parser = load_ini(
        INVOICE_SOURCES_INI
    )

    result = []

    for section in parser.sections():

        if not section.startswith(
            "Source."
        ):
            continue

        enabled = parser.get(
            section,
            "Enabled",
            fallback="1",
        ).strip()

        if not is_enabled_flag(enabled):
            continue

        raw_path = parser.get(
            section,
            "Path",
            fallback="",
        ).strip()

        if not raw_path:
            continue

        result.append(
            RegisteredSource(
                source_key=section[7:],
                path=Path(raw_path),
                owner=parser.get(
                    section,
                    "Owner",
                    fallback="",
                ).strip(),
                store_code=parser.get(
                    section,
                    "StoreCode",
                    fallback="",
                ).strip(),
                carrier_code=parser.get(
                    section,
                    "CarrierCode",
                    fallback="",
                ).strip().upper(),
            )
        )

    return result


def load_registered_sources() -> list[Path]:
    """兼容旧调用：只返回已启用扫描源路径。"""

    return [
        entry.path
        for entry in load_source_entries()
    ]


# =========================================================
# 收集 Excel
# =========================================================

def resolve_invoice_scan_directories(
    *,
    source_directory: Path | None,
    extra_directories: list[Path],
    extra_only: bool,
) -> list[Path]:
    """原始扫描只用当前物流商扫描源；临时追加扫描只用手动选择的目录。"""

    if extra_only:
        return list(extra_directories)
    if source_directory is None:
        return []
    return [source_directory]


def collect_invoice_files(
    directories: list[Path],
) -> list[Path]:

    files = []

    seen = set()

    for directory in directories:

        if not directory.exists():
            continue

        if not directory.is_dir():
            continue

        for path in directory.rglob(
            "*"
        ):

            if not path.is_file():
                continue

            if path.name.startswith(
                "~$"
            ):
                continue

            if is_excel_junk_file(path):
                continue

            if (
                path.suffix.lower()
                not in ALLOWED_EXTENSIONS
            ):
                continue

            relative_parts = set(
                path.relative_to(
                    directory
                ).parts[:-1]
            )

            if (
                relative_parts
                & EXCLUDE_DIR_NAMES
            ):
                continue

            resolved = path.resolve()

            if resolved in seen:
                continue

            seen.add(resolved)

            files.append(path)

    return sorted(
        files,
        key=lambda x:
        str(x).lower(),
    )


# =========================================================
# 读取 _SystemMeta
# =========================================================

def read_system_meta(
    path: Path,
) -> tuple[
    dict | None,
    str | None,
]:

    try:

        workbook = load_workbook(
            path,
            read_only=False,
            data_only=False,
            keep_links=False,
        )

    except Exception as exc:

        return None, (
            f"{path.name}："
            f"Excel 无法读取：{exc}"
        )

    try:

        if (
            "_SystemMeta"
            not in workbook.sheetnames
        ):

            return None, (
                f"{path.name}："
                "缺少 _SystemMeta"
            )

        ws = workbook[
            "_SystemMeta"
        ]

        if (
            ws.sheet_state
            != "veryHidden"
        ):

            return None, (
                f"{path.name}："
                "_SystemMeta 必须为 veryHidden"
            )

        values = {}

        for row in ws.iter_rows(
            min_col=1,
            max_col=2,
            values_only=True,
        ):

            key = row[0]
            value = row[1]

            if key is None:
                continue

            key = str(
                key
            ).strip()

            if not key:
                continue

            values[key] = (
                ""
                if value is None
                else str(value).strip()
            )

        missing = [
            field
            for field in META_FIELDS
            if not values.get(field)
        ]

        if missing:

            return None, (
                f"{path.name}："
                "_SystemMeta 缺失字段："
                + ", ".join(missing)
            )

        return values, None

    finally:

        workbook.close()


def _keep_latest_generation(
    staged: list[tuple[Path, dict]],
) -> tuple[list[tuple[Path, dict]], list[str]]:
    """
    同一 SourceID 若被重新生成过，未合并里常会留下旧文件。
    GeneratedAt 可以不同；只保留最新一批，旧文件跳过。
    """

    latest: dict[str, str] = {}

    for _path, meta in staged:
        source_id = str(meta.get("SourceID") or "")
        generated_at = str(meta.get("GeneratedAt") or "")
        current = latest.get(source_id)
        if current is None or generated_at > current:
            latest[source_id] = generated_at

    kept = []
    skipped: dict[str, list[str]] = {}

    for path, meta in staged:
        source_id = str(meta.get("SourceID") or "")
        generated_at = str(meta.get("GeneratedAt") or "")
        if generated_at != latest.get(source_id, generated_at):
            skipped.setdefault(source_id, []).append(path.name)
            continue
        kept.append((path, meta))

    warnings = []
    for source_id, names in skipped.items():
        warnings.append(
            f"SourceID {source_id} 未合并里混有更早生成的旧发票，"
            f"已跳过 {len(names)} 个，只保留最新批次 "
            f"{latest.get(source_id, '')}："
            + "、".join(names)
        )

    return kept, warnings


# =========================================================
# 主扫描
# =========================================================

def scan_original_invoices(
    date_id: str,
    directories: list[Path],
) -> ScanResult:
    """
    只扫描未合并发票本身。
    不读取询价明细，不与询价中心对齐。
    """

    errors = []
    warnings = []

    try:

        known_carriers = (
            load_known_carriers()
        )

    except Exception as exc:

        return ScanResult(
            passed=False,
            invoice_count=0,
            carton_count=0,
            source_count=0,
            warehouse_count=0,
            carrier_count=0,
            records=[],
            errors=[str(exc)],
            warnings=[],
        )

    files = collect_invoice_files(
        directories
    )

    if not files:

        errors.append(
            "没有扫描到任何发票 .xlsx"
        )

    records = []

    found_keys = set()
    global_cartons = set()

    source_consistency = {}
    source_mismatch_reported = set()
    fba_warehouse = {}

    staged: list[tuple[Path, dict]] = []

    for path in files:

        meta, meta_error = (
            read_system_meta(path)
        )

        if meta_error:

            errors.append(
                meta_error
            )

            continue

        if (
            meta["DateID"]
            != date_id
        ):

            continue

        staged.append((path, meta))

    kept, generation_warnings = _keep_latest_generation(
        staged
    )
    warnings.extend(generation_warnings)

    for path, meta in kept:

        carrier_code = (
            meta["CarrierCode"]
            .strip()
            .upper()
        )

        if (
            carrier_code
            not in known_carriers
        ):

            errors.append(
                f"{path.name}："
                f"未知 CarrierCode："
                f"{carrier_code}"
            )

        if (
            meta["MetaSchemaVersion"]
            != "1.0"
        ):

            errors.append(
                f"{path.name}："
                "MetaSchemaVersion "
                "必须为 1.0"
            )

        source_id = meta[
            "SourceID"
        ]

        carton = meta[
            "CartonNumber"
        ]

        fba_batch = meta[
            "FBABatch"
        ]

        warehouse = meta[
            "WarehouseCode"
        ]

        source_error = validate_source_id(
            source_id
        )

        if source_error:

            errors.append(
                f"{path.name}："
                f"{source_error}"
            )

        carton_error = validate_carton_number(
            fba_batch,
            carton,
        )

        if carton_error:

            errors.append(
                f"{path.name}："
                f"{carton_error}"
            )

        # ======================================
        # Machine Key
        # ======================================

        key = machine_key(
            source_id,
            carton,
        )

        if key in found_keys:

            errors.append(
                "重复发票 Key："
                f"{source_id} + {carton}"
            )

        else:
            found_keys.add(key)

        # CartonNumber 全局唯一
        if carton in global_cartons:

            errors.append(
                "CartonNumber 全局重复："
                f"{carton}"
            )

        else:
            global_cartons.add(
                carton
            )

        # ======================================
        # SourceID 一致性
        # ======================================

        consistency_value = (
            meta["StoreCode"],
            meta["PlanID"],
            meta["DateID"],
            carrier_code,
        )

        previous = (
            source_consistency.get(
                source_id
            )
        )

        if previous is None:

            source_consistency[
                source_id
            ] = consistency_value

        elif (
            previous != consistency_value
            and
            source_id not in source_mismatch_reported
        ):

            errors.append(
                f"SourceID {source_id} "
                "内部元数据不一致："
                "同一 SourceID 的店铺/计划/日期/物流商必须相同。"
                f"文件 {path.name} 与同批其他发票不一致。"
            )
            source_mismatch_reported.add(
                source_id
            )

        # ======================================
        # FBA → Warehouse
        # ======================================

        previous_wh = (
            fba_warehouse.get(
                fba_batch
            )
        )

        if (
            previous_wh
            and
            previous_wh != warehouse
        ):

            errors.append(
                f"FBABatch {fba_batch} "
                "同时属于多个 Warehouse："
                f"{previous_wh} / {warehouse}"
            )

        else:

            fba_warehouse[
                fba_batch
            ] = warehouse

        # ======================================
        # SHA256
        # ======================================

        try:

            file_hash = (
                sha256_file(path)
            )

        except Exception as exc:

            errors.append(
                f"{path.name}："
                f"SHA256 失败：{exc}"
            )

            continue

        records.append(
            InvoiceRecord(
                path=str(path),
                sha256=file_hash,

                carrier_code=
                    carrier_code,

                carrier_name=
                    meta[
                        "CarrierName"
                    ],

                template_version=
                    meta[
                        "TemplateVersion"
                    ],

                date_id=
                    meta["DateID"],

                generated_at=
                    meta["GeneratedAt"],

                store_code=
                    meta["StoreCode"],

                plan_id=
                    meta["PlanID"],

                source_id=
                    source_id,

                warehouse_code=
                    warehouse,

                fba_batch=
                    fba_batch,

                carton_number=
                    carton,
            )
        )

    return ScanResult(
        passed=not errors,

        invoice_count=len(
            records
        ),

        carton_count=len(
            global_cartons
        ),

        source_count=len({
            record.source_id
            for record in records
        }),

        warehouse_count=len({
            record.warehouse_code
            for record in records
        }),

        carrier_count=len({
            record.carrier_code
            for record in records
        }),

        records=records,

        errors=errors,
        warnings=warnings,
    )
 