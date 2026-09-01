from __future__ import annotations

import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from app.shipment_tracking.domain.parsed_plan import (
    ParsedCarton,
    ParsedCartonItem,
    ParsedFBA,
    ParsedTrackingPlan,
)

PLAN_SHEET_NAME = re.compile(r"^.+-发货规划$")
STORE_TO_US_PLAN = re.compile(r"^美(\d+)$")
FBA_ID_PATTERN = re.compile(r"^FBA[0-9A-Za-z-]+$", re.IGNORECASE)


def parse_shipping_plan(
    workbook_path,
    *,
    store_code: str | None = None,
    planning_sheet: str | None = None,
) -> ParsedTrackingPlan:
    """
    只读发货规划矩阵，输出 FBA → Carton → SKU → Qty。
    不扫描目录、不算发票箱号、不聚合、不写库。
    """

    path = Path(workbook_path)
    workbook = load_workbook(path, data_only=True)
    try:
        worksheet = _choose_planning_sheet(
            workbook,
            store_code=store_code,
            planning_sheet=planning_sheet,
        )
        fba_columns = _read_fba_columns(worksheet)
        if not fba_columns:
            return ParsedTrackingPlan(fbas=())
        items_by_column = _read_sku_items(worksheet, fba_columns)
        return _build_plan(fba_columns, items_by_column)
    finally:
        workbook.close()


def _choose_planning_sheet(workbook, *, store_code, planning_sheet):
    names = list(workbook.sheetnames)
    if planning_sheet:
        if planning_sheet not in names:
            raise ValueError(f"没有工作表【{planning_sheet}】")
        return workbook[planning_sheet]

    mapped = _sheet_name_for_store(store_code)
    if mapped and mapped in names:
        return workbook[mapped]

    candidates = [name for name in names if PLAN_SHEET_NAME.match(name)]
    if len(candidates) == 1:
        return workbook[candidates[0]]
    if not candidates:
        raise ValueError("没有发货规划工作表")
    raise ValueError("多个发货规划工作表，需要 store_code 或 planning_sheet")


def _sheet_name_for_store(store_code: str | None) -> str | None:
    if not store_code:
        return None
    match = STORE_TO_US_PLAN.fullmatch(store_code.strip())
    if not match:
        return None
    return f"美国{int(match.group(1))}号-发货规划"


def _read_fba_columns(worksheet) -> list[dict]:
    fba_row = _find_row_label(worksheet, "FBA批号")
    warehouse_row = _find_row_label(worksheet, "仓库")
    length_row = _find_row_label(worksheet, "长(cm)")
    width_row = _find_row_label(worksheet, "宽(cm)")
    height_row = _find_row_label(worksheet, "高(cm)")
    weight_row = _find_row_label(worksheet, "重量(KG)")
    if fba_row is None:
        return []

    header_row = _find_sku_header_row(worksheet)
    columns = []
    max_col = worksheet.max_column or 14
    for col in range(1, max_col + 1):
        fba_id = _text(worksheet.cell(fba_row, col).value)
        if not FBA_ID_PATTERN.fullmatch(fba_id):
            continue
        carton_no = None
        if header_row is not None:
            carton_no = _text(worksheet.cell(header_row, col).value) or None
        columns.append(
            {
                "col": col,
                "fba_id": fba_id,
                "destination_fc": (
                    _text(worksheet.cell(warehouse_row, col).value)
                    if warehouse_row
                    else ""
                )
                or None,
                "carton_no": carton_no,
                "length_cm": _decimal(
                    worksheet.cell(length_row, col).value if length_row else None
                ),
                "width_cm": _decimal(
                    worksheet.cell(width_row, col).value if width_row else None
                ),
                "height_cm": _decimal(
                    worksheet.cell(height_row, col).value if height_row else None
                ),
                "weight_kg": _decimal(
                    worksheet.cell(weight_row, col).value if weight_row else None
                ),
            }
        )
    return columns


def _read_sku_items(worksheet, fba_columns: list[dict]) -> dict[int, list[ParsedCartonItem]]:
    header_row = _find_sku_header_row(worksheet)
    if header_row is None:
        return {}
    start_row = header_row + 1
    max_row = worksheet.max_row or start_row
    items_by_column: dict[int, list[ParsedCartonItem]] = defaultdict(list)

    for row in range(start_row, max_row + 1):
        sku = _text(worksheet.cell(row, 6).value)
        if not sku:
            continue
        asin = _text(worksheet.cell(row, 5).value) or None
        en = _text(worksheet.cell(row, 1).value)
        cn = _text(worksheet.cell(row, 8).value)
        product_name = en or cn or None
        for column in fba_columns:
            qty = _quantity(worksheet.cell(row, column["col"]).value)
            if qty <= 0:
                continue
            items_by_column[column["col"]].append(
                ParsedCartonItem(
                    sku=sku,
                    quantity=qty,
                    product_name=product_name,
                    asin=asin,
                )
            )
    return items_by_column


def _build_plan(
    fba_columns: list[dict],
    items_by_column: dict[int, list[ParsedCartonItem]],
) -> ParsedTrackingPlan:
    grouped: dict[str, list[ParsedCarton]] = defaultdict(list)
    destinations: dict[str, str | None] = {}

    for column in fba_columns:
        items = tuple(items_by_column.get(column["col"], ()))
        if not items:
            continue
        fba_id = column["fba_id"]
        destinations.setdefault(fba_id, column["destination_fc"])
        grouped[fba_id].append(
            ParsedCarton(
                fba_id=fba_id,
                destination_fc=column["destination_fc"],
                carton_no=column["carton_no"],
                weight_kg=column["weight_kg"],
                length_cm=column["length_cm"],
                width_cm=column["width_cm"],
                height_cm=column["height_cm"],
                items=items,
            )
        )

    fbas = []
    for fba_id, cartons in grouped.items():
        fbas.append(
            ParsedFBA(
                fba_id=fba_id,
                destination_fc=destinations.get(fba_id),
                cartons=tuple(cartons),
            )
        )
    return ParsedTrackingPlan(fbas=tuple(fbas))


def _find_row_label(worksheet, label: str) -> int | None:
    target = label.strip()
    max_row = min(worksheet.max_row or 1, 20)
    max_col = min(worksheet.max_column or 1, 20)
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            value = _text(worksheet.cell(row, col).value)
            if value == target or value.lstrip("*") == target:
                return row
    return None


def _find_sku_header_row(worksheet) -> int | None:
    max_row = min(worksheet.max_row or 1, 30)
    for row in range(1, max_row + 1):
        if _text(worksheet.cell(row, 6).value).upper() == "SKU":
            return row
    return None


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _quantity(value) -> int:
    if value in (None, ""):
        return 0
    try:
        number = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        return 0
    if number <= 0:
        return 0
    as_int = int(number)
    if number != as_int:
        return int(number.to_integral_value())
    return as_int


def _decimal(value) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        number = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        return None
    return number
